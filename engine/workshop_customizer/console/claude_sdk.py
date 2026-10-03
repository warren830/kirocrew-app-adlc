"""Console module ``claude_sdk`` — 从模板生成：Claude Agent SDK. An agent described by a spec — prompt, model, turns,
tools, skills, knowledge bases — and no code: the console writes its container (Launchpad's 方式A) and deploys it
through the console's own deploy pipeline (``deploy.start_deploy``, source ``dockerfile``: CodeBuild on ARM, the
runtime's own ECR repository that scans on push, the scan gate, CreateAgentRuntime or a new version, a smoke call).

The build context (:func:`generate`, zipped by :func:`bundle`, the same spec making the same zip):

* ``Dockerfile`` — linux/arm64 ``python:3.12-slim-trixie`` from ECR Public's mirror of the Docker official images
  (CodeBuild is never rate-limited by Docker Hub), the base image's security updates (the scan gate reads this
  image), Node 22 (nodejs.org, its SHA-256 checked) with the Claude Code CLI ``@anthropic-ai/claude-code`` pinned to the version the SDK
  release bundles, the SDK and bedrock-agentcore (boto3 checked ≥ 1.43.90 at build time), a non-root user ``agent``
  (uid 1001, ``HOME=/app``), ``CLAUDE_CODE_USE_BEDROCK=1``, port 8080; Debian's ``fonts-wqy-microhei`` and a
  matplotlibrc (``MPLCONFIGDIR=/app/.matplotlib``, Agg) whose default font has Chinese glyphs, its font cache built
  at build time and the font checked there.
* ``main.py`` — a ``bedrock_agentcore.runtime.BedrockAgentCoreApp`` whose entrypoint answers ``{"prompt"}`` with
  ``claude_agent_sdk.query()``: the spec's system prompt, model (a Bedrock inference profile id), ``max_turns``,
  exactly the built-in tools the spec allows (``tools``, so nothing else is even offered, plus ``Skill`` when there
  are skills), ``permission_mode="dontAsk"`` (what is not pre-approved is refused, never prompted), the skills by
  name (the SDK's ``skills`` option), the knowledge bases as an in-process MCP server ``kb`` with one tool
  ``retrieve`` (``mcp__kb__retrieve``: bedrock-agent-runtime Retrieve, ``managedSearchConfiguration`` on a managed
  KB), ``strict_mcp_config`` and ``verbatim_prompts`` (a user's ``@/path`` or ``/command`` is text, not an action).
  It answers ``text/plain`` with a 〔工具〕 line of the turn's tool calls and a 〔文件〕 line of the files it made
  (the console's chat shows a runtime's body as it is), ``{"format": "json"}`` with ``{"result", "tools", "skills",
  "skillsLoaded", "artifacts", "turns", "costUsd", ...}``, and ``{"action": "workspace"}`` with the workspace's files.
* ``requirements.txt`` (with openpyxl, pypdf, python-docx, python-pptx, Pillow and matplotlib: the agent's Bash writes
  a spreadsheet, a deck or a chart with them, the answer reads them back), ``agent_spec.json`` (the spec, with the skill
  versions used) and each skill of the workspace's skill library as ``.claude/skills/<name>/SKILL.md``.

**A call's own skills** (template 2, the Skill Lab's arms): the payload may carry ``skills`` (the names this call may
load; ``[]`` none: the arm without the skill) and ``skillOverrides`` (``[{"name", "uri": "s3://<console bucket>/skills/
<name>/<version>/"}]``: that SKILL.md and its folder read from S3 for this call, so a library version or a training's
candidate is tried without a new image; the runtime's role reads ``skills/*`` once the Skill Lab grants it). Such a
call runs on its own: a working directory ``calls/<id>`` in the workspace (what it writes is its own), a Claude Code
config of its own whose ``skills/`` holds exactly the call's skills (the image's copied in, the overrides fetched),
``setting_sources=["user"]`` (no project skill of the workspace or the image joins in), ``Skill`` offered only when
there is a skill, a fresh Claude Code session that is not resumed nor remembered. A refusal is answered 400 (a payload
to fix) or 503 (an override S3 would not give: a grant made seconds ago), each with a ``call refused`` log line.

**A call's input files** (template 3, a Skill Lab task's): ``files`` — ``[{"path", "uri": "s3://<console bucket>/
skill-lab/assets/<sha256>", "sha256", "bytes"}]``, at most 32, 25 MiB each, 100 MiB in all, each ``path`` a plain
relative path (4 levels, no hidden names) — makes a call of its own too. Each file is read from S3 into the call's
directory before the prompt, its size and SHA-256 checked (a mismatch or a missing object is a 400, a read S3 refuses a
503), and the prompt ends with a 〔输入文件〕 line naming them; the directory's manifest is taken after, so an input the
turn leaves alone is not one of its files and one it rewrites is ``changed``. The answer's ``call.inputs`` lists them
(``[]`` when there are none: how the console tells a template-3 image from an older one).

**The files a turn made** (``artifacts``, every call): the regular files under its directory (the session's workspace
for a conversation's turn, without ``calls/``; Claude Code's own and link targets left out) whose size or mtime
changed during the turn, each with ``path``, ``bytes``, ``type``, ``status`` (created / changed), a text view the judge
reads (a spreadsheet's sheets as numbered rows of cells — formulas as written, a cached value after ``→``, a text cell
that looks like a number in quotes, the frozen pane and a bold header row noted; a CSV or text file as it is; a Word
document's headings, paragraphs and tables; a PDF's pages; a presentation's slides — titles, text by level, tables, a
chart's type, title, categories and series, pictures, notes), a table preview for xlsx / csv / docx / pptx tables, and
the bytes in base64 when the file is at most 48 KB (160 KB per answer); at most 20 files, 6000 characters each, 24000
in all. **Pictures** (template 3): an image (png, jpg, gif, webp, bmp, tiff) gets ``image`` — its format and size, a
``preview`` for the judge (the image itself when it is a Converse format of at most 1568 px and 1 MB, else Pillow's
PNG — JPEG for a photo or when the PNG is over 1 MB — at most 1568 px, EXIF orientation applied) and a 320 px ``thumb``
for the console's page; a Word document's or a workbook's media and a presentation's slide pictures (4 a document)
get the same under ``images``; at most 8 previews, 4 MB, an answer, the rest with a ``note`` (an EMF Pillow cannot
read: an ``error``).

A conversation is a runtime session: the first turn starts a Claude Code session whose id is derived from the runtime
session id, the next ones ``resume`` it. ``CLAUDE_CONFIG_DIR`` — the Claude Code session transcripts, the skills
(copied in at startup) — lives in the agent's workspace, which is the runtime's session storage when the deployment
turns it on (``AGENT_WORKSPACE``, set from ``sessionStorage``): the files the agent writes and the conversation then
survive a stopped session (AgentCore resets session storage on a new runtime version and after 14 idle days).

Records (``claude_sdk_agents`` in the console store) keep each agent's spec and its deployments; a runtime named
``adlc_probe_<module>_*`` deploys with its module's probe names (``studio.names_for``).

Local checks on 2026-10-02 (claude-agent-sdk 0.2.163, its bundled CLI 2.1.286, Bedrock in us-west-2):

* ``tools=[]`` turns off every built-in tool **including Skill**: an agent with skills must list ``Skill`` in
  ``tools``, then the model calls ``Skill {"skill": "<name>"}``. ``CLAUDE_CONFIG_DIR/skills/<name>/SKILL.md`` is found
  as a user skill (the SDK's ``skills`` option sets ``setting_sources`` to user and project).
* ``permission_mode="dontAsk"`` still runs read-only Bash commands (``ls``) that are not in ``allowed_tools``: the
  ``tools`` list, not the permission mode, is what keeps a tool away from an agent.
* ``resume`` of a session whose transcript is missing raises ``ResultError: No conversation found with session ID``;
  ``session_id=<uuid5>`` on the first turn and ``resume`` afterwards keep one conversation per runtime session.
* With ``verbatim_prompts`` Claude Sonnet 5.5 still loads the skill on the first call (the listing is in the Skill
  tool); ``us.anthropic.claude-sonnet-5-5`` and ``us.anthropic.claude-haiku-4-5-20251001-v1:0`` work as ``model``.
* The SDK's platform wheels bundle the CLI (225 MB): the image keeps one copy, the npm one on ``PATH``
  (``cli_path``). Claude Code on Bedrock has no WebSearch (no server-side web search there); WebFetch fetches from
  the runtime itself.

Live 2026-10-02 (us-west-2, ``adlc_probe_claude_64cde0``: Claude Sonnet 5.5, Read / Write / Glob, the library's
``loyalty-reply-style`` v0001, a managed KB, session storage at /mnt/workspace, lifecycle 600 /
7200 s; the 10 KB build context became a 252 MB image):

* **Deploy, 3 min** (the deploy module's docstring has the stages); the default CRITICAL scan gate stopped the first
  attempt — CRITICAL 4 on the bookworm base (perl ×3, python3.11 pulled in by NodeSource's nodejs) — and
  ``scanBlock: []`` deployed it. Since then the base is Debian 13 (``python:3.12-slim-trixie``) and Node comes from
  nodejs.org (SHA-256 checked): the same spec scanned CRITICAL 0, HIGH 2, MEDIUM 1 and deployed with the default gate
  (``adlc_probe_claude_a56879``: build 125 s, scan 127 s, smoke from the KB 16.9 s). The smoke call (the spec's sample question) used both tools, 17.8 s cold.
* **Chat through the console's runtime route** (``{"prompt", "actorId"}``, the body shown as it is): turn 1 12.1 s,
  ``retrieve`` then ``Skill(loyalty-reply-style)``, the answer in the skill's form (the conclusion first, the rule
  named, the signature line) with the current rule over the KB's superseded 2023 one; turn 2 5.1 s, ``Write``
  notes.md into /mnt/workspace; StopRuntimeSession (11.5 s); turn 3 2.6 s on a new microVM (a second ``ready`` log
  line) resumed the same Claude Code session from the session storage and recalled turns 1 and 2. $0.01–0.02 a turn.
* The npm CLI (``/usr/bin/claude``, 2.1.286) ran; with pinned Bedrock model ids nothing asked for
  ListInferenceProfiles and the CLI wrote nothing to stderr. Uvicorn logs ``Invalid HTTP request received`` once per
  microVM start. A container runtime's log streams are ``<date>/[runtime-logs]<instance id>``, one per microVM (not
  per session id, as a code runtime's are).
* A new version (another prompt) kept the session storage — GetAgentRuntime's ``filesystemConfigurations`` passed
  back as they are — and AgentCore gave the old session an empty one, as it documents.

Local checks on 2026-10-02 for a call's own skills (the same SDK and CLI): Claude Code finds project skills in the
working directory's ``.claude/skills`` **and in its parents'** (an ancestor's skill was listed); a user skill
(``CLAUDE_CONFIG_DIR/skills``) beats a project skill of the same name; ``setting_sources=["project"]`` drops the user
skills; ``setting_sources=["user"]`` with the call's own ``CLAUDE_CONFIG_DIR`` lists only that config's skills (asked
for one in the working directory or a parent, the model answered that only the call's skill exists) — hence a call's
own config rather than project skills. ``skills=[names]`` is a filter (the Skill tool refuses an unlisted one, the init
message still lists every skill found); with ``skills=[]`` and Skill offered the model still tried ``Skill(alpha)`` and
was refused, so the arm without skills is not offered Skill at all. The template's ``main.py`` on the real SDK with
the sample ``points-report-xlsx`` (S3 faked): with the skill 10 s, ``Skill`` then ``Bash`` (python3 + openpyxl), the
sheet as the skill says; without it ``会员积分报表.csv`` with a ``会员编号`` header, a fourth column and totals written
as numbers.

Live 2026-10-02 (``adlc_probe_artifacts_dc2877``: Claude Sonnet 5.5, Read / Write / Edit / Bash / Glob, the sample
``points-report-xlsx``, session storage, lifecycle 600 / 3600 s):

* **Deploy, 2.6 min**: build 92 s, scan 22 s — HIGH 2, MEDIUM 1, so the default CRITICAL gate passed with the file
  readers (and lxml, python-docx's) in the image — runtime 14 s, ready 6 s, smoke 18.7 s.
* AgentCore keeps a container's 4xx and 5xx bodies to itself: ``skills: ["nope"]`` came back as ``RuntimeClientError:
  Received error (400) from runtime``, an override read before the grant as ``(503)`` (the log: AccessDenied on
  ListObjectsV2 for the runtime's role) — the reason is only in the log, hence a ``call refused`` line for each.
  StopRuntimeSession took 11 s.
* **Twelve calls of the Skill Lab** (the sample's 6 test tasks, with and without the skill, 4 at a time, 60 s in
  all): with it 5.5–6.6 s and 4 turns (``Skill``, ``Bash``), $0.011–0.013, ``points-report.xlsx`` each time; without
  it 6.5–10 s, ``积分明细表.csv``, ``会员积分报表.xlsx``, ``会员积分.xlsx`` …, each with other sheets or columns.
* **Publishing** to it (the Skill Lab rebuilds the agent with the trained version: ``deploy_agent(runtimeId=…)``):
  build 82 s, scan 11 s, UpdateAgentRuntime 3.5 s, ready 6 s, smoke 16 s; a conversation's turn on version 2 loaded
  the new version from the image and its text answer ended ``〔文件〕points-report.xlsx（4.9 KB）``.

Local checks on 2026-10-02 for template 3, its ``main.py`` with the real python-pptx 1.0.2, Pillow 12.3.0, matplotlib
3.11.2, openpyxl and python-docx (the SDK stubbed): a deck with a title and bulleted slide, notes, a table, a bar
chart and a picture came back as its slides' text (``[图表 BAR_CLUSTERED] 标题「积分图」；类别 M1、M2；系列「本期积分」10, 400``)
with the slide's picture previewed; a docx's and an xlsx's media pictures previewed; an 800×450 chart passed as its
own preview (9.8 KB), a 4000×3000 PNG became a 1568×1176 preview, a 3000×2000 JPEG a 1568×1045 JPEG; the input file in
the directory before the turn was not reported. matplotlib indexes Debian's ``wqy-microhei.ttc`` as "WenQuanYi Micro
Hei" (face 0) and draws 会员本期积分 with it, no missing-glyph warning.

Live 2026-10-02 (``adlc_probe_inputs_bf820b``: template 3, Claude Sonnet 5.5, Read / Write / Edit / Bash / Glob, the
sample ``points-chart-png``, session storage, lifecycle 600 / 3600 s):

* **Deploy, 2.7 min**: build 103 s (CodeBuild 1.5 min, build phase 60 s), scan 11 s — HIGH 2 (gcc-14, zlib), MEDIUM 1
  (dash): the base image's, so the default CRITICAL gate passed — runtime 14 s, ready 6 s, smoke 24 s (the invoke 13 s).
  The image is 317 MB in ECR (252 MB before python-pptx, Pillow, matplotlib with numpy, and the font). The account's
  registry scans BASIC: OS packages only, so the pip libraries are not in the gate's findings at all (an enhanced,
  Inspector scan would read them).
* **The Skill Lab's calls with input files** (8, 4 at a time): the runtime's ``turn`` log line names each call's
  ``inputs`` (``members.csv``; ``members.csv``, ``supplement.csv``; ``members.xlsx``) apart from its ``files``; with the
  skill 7.6–9.3 s (``Skill``, two ``Bash``), without it 15.8–21.4 s (4–6 ``Bash``, a ``Read``). Every chart drew its
  Chinese title with the image's font; one arm without the skill still told the user "这台机器没有中文字体" and drew English
  labels — there is no fontconfig, so ``fc-list`` finds nothing, while matplotlib has the font.
* **A deck** (one call): ``chart.png`` 1050×600 and a two-slide ``report.pptx`` (a table, the chart as a picture), 17 s,
  $0.029: the slides' text and the slide picture's preview came back, and the judge saw both pictures.
"""
from __future__ import annotations

import base64
import hashlib
import json
import re
from typing import Any, Mapping, Sequence

from ..direct.aws import client
from . import deploy, skills_lab, studio
from .common import now as _now, safe
from .studio import _clean, _line, bundle, names_for, py_str, py_value

SDK = "claude-agent-sdk==0.2.163"
#: The Claude Code CLI version claude-agent-sdk 0.2.163 bundles (``_cli_version.py``): npm installs the same one.
CLAUDE_CODE = "2.1.286"
AGENTCORE = "bedrock-agentcore==1.24.0"
#: Retrieve on a managed KB takes ``managedSearchConfiguration``, which older botocore models refuse client-side.
BOTO = ("boto3>=1.43.90,<2", "botocore>=1.43.90,<2")
BASE_IMAGE = "public.ecr.aws/docker/library/python:3.12-slim-trixie"
NODE_MAJOR = 22
#: Node from nodejs.org, checked against its published SHA-256: NodeSource's nodejs package pulled python3.11 into
#: the image, and with Debian 12's perl it made the default CRITICAL scan gate refuse the template's own image (live).
NODE_VERSION = "22.23.3"
NODE_SHA256 = "a44aeb94849a299b22df10b9e622ec2f605c2183501bc40590705131de7c740f"  # node-v22.23.3-linux-arm64.tar.xz
#: The libraries that read the files a turn makes (xlsx, pdf, docx, pptx; Pillow for an image's preview) for its
#: answer's ``artifacts``; the agent's own Bash finds them too (``python3`` with openpyxl writes a spreadsheet).
FILE_READERS = ("openpyxl==3.1.5", "pypdf==6.19.0", "python-docx==1.2.0", "python-pptx==1.0.2", "pillow==12.3.0")
#: The agent's Bash draws charts: matplotlib (Agg; a font with Chinese, ``CHART_FONT``, the default).
CHARTS = ("matplotlib==3.11.2",)
#: Debian's WenQuanYi Micro Hei (4.6 MB): Latin and CJK glyphs, so a chart's Chinese labels are not boxes.
CHART_FONT, CHART_FONT_PACKAGE = "WenQuanYi Micro Hei", "fonts-wqy-microhei"
#: 2: a call may bring its own skills (``skills`` / ``skillOverrides``) and the answer lists the files the turn made.
#: 3: a call may bring its input files (``files``), and the answer brings pictures (an image's, a document's) and a
#: presentation's slides for the judge.
TEMPLATE = 3
FORMAT = f"adlc-claude-agent/{TEMPLATE}"
COLLECTION = "claude_sdk_agents"
KB_SERVER, KB_TOOL_NAME = "kb", "retrieve"
KB_TOOL = f"mcp__{KB_SERVER}__{KB_TOOL_NAME}"

DEFAULT_MODEL = "us.anthropic.claude-sonnet-5-5"
DEFAULT_FAST_MODEL = "us.anthropic.claude-haiku-4-5-20251001-v1:0"
MODELS = (
    ("us.anthropic.claude-sonnet-5-5", "Claude Sonnet 5.5"),
    ("us.anthropic.claude-haiku-4-5-20251001-v1:0", "Claude Haiku 4.5（快、便宜）"),
    ("us.anthropic.claude-opus-5-5", "Claude Opus 5.5"),
    ("us.anthropic.claude-sonnet-5", "Claude Sonnet 5"),
    ("us.anthropic.claude-sonnet-4-5-20250929-v1:0", "Claude Sonnet 4.5"),
)
#: Claude Code's built-in tools an agent may be given (WebSearch is not on Bedrock; subagents are off).
BUILTIN_TOOLS: dict[str, str] = {
    "Read": "读文件", "Write": "写文件", "Edit": "改文件", "Glob": "按名字找文件", "Grep": "按内容搜文件",
    "Bash": "运行命令（在 Runtime 的容器里）", "WebFetch": "读网页（从 Runtime 访问外网）",
}
EFFORTS = ("low", "medium", "high", "xhigh", "max")
MAX_PROMPT, MAX_SKILLS, MAX_KBS, MAX_SAMPLE = 20_000, 10, 5, 2000
TURNS_MIN, TURNS_MAX, DEFAULT_TURNS = 1, 50, 12
TOP_K_MIN, TOP_K_MAX, DEFAULT_TOP_K = 1, 20, 5
MODEL_ID = re.compile(r"^(?:[a-z0-9][a-z0-9.:_\-/]{2,199}|arn:aws[a-z-]*:bedrock:[a-z0-9-]+:\d{12}:[a-z-]+/[A-Za-z0-9._:/-]{1,200})$")
SKILL_REF = re.compile(r"^(?P<name>[a-z0-9](?:[a-z0-9-]{0,62}[a-z0-9])?)(?:@(?P<version>v\d{4}))?$")
DEFAULT_PROMPT = "你是一个乐于助人的助手。用中文简洁、准确地回答；不知道就直说。"
DEFAULT_SAMPLE = "你好！请用一句话介绍你自己。"

SAMPLE_SPEC: dict[str, Any] = {
    "description": "拾光家居会员积分客服：知识库检索规则，技能规范回复（示例）",
    "systemPrompt": "你是拾光家居的会员积分客服。积分规则（获取、有效期、推荐奖励、补偿）先用知识库检索（retrieve），按检索到的原文回答，"
                    "不要凭记忆回答；回复按技能 loyalty-reply-style 的规范来写。会员要你记下什么，就用 Write 写到工作目录的 notes.md。用中文。",
    "model": DEFAULT_MODEL, "fastModel": DEFAULT_FAST_MODEL, "maxTurns": DEFAULT_TURNS, "tools": ["Read", "Write", "Glob"],
    "skills": ["loyalty-reply-style"], "knowledgeBases": [], "sample": "推荐好友注册下单后，我什么时候能拿到推荐积分？",
}


class ClaudeSdkError(ValueError):
    """A spec or request the caller must fix (the route answers 400)."""


# -- the spec ------------------------------------------------------------------------------------------------------------

def _int(value: Any, default: int, lo: int, hi: int, what: str) -> int:
    if value is None or value == "":
        return default
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ClaudeSdkError(f"{what}：要一个整数（{lo} 到 {hi}）") from exc
    if isinstance(value, bool) or number != float(value) or not lo <= number <= hi:
        raise ClaudeSdkError(f"{what}：{lo} 到 {hi} 之间的整数")
    return number


def check_spec(body: Mapping[str, Any]) -> dict[str, Any]:
    """The agent the form describes, typed and checked (every refusal in Chinese, naming the field)."""
    if not isinstance(body, Mapping):
        raise ClaudeSdkError("spec：一个对象")
    name = str(body.get("name") or "").strip()
    if not deploy.NAME.match(name):
        raise ClaudeSdkError("名称：字母开头，字母、数字或下划线，最多 48 个（它就是 Runtime 的名称）")
    if name.lower().startswith("harness_"):
        raise ClaudeSdkError("名称不能以 harness_ 开头：那是 Harness 的 Runtime")
    prompt = _clean(body.get("systemPrompt")).strip()
    if not prompt:
        raise ClaudeSdkError("系统 Prompt 是空的：写清楚它是谁、做什么、怎么回答")
    if len(prompt) > MAX_PROMPT:
        raise ClaudeSdkError(f"系统 Prompt 最多 {MAX_PROMPT} 个字符（现在 {len(prompt)}）")
    model = str(body.get("model") or DEFAULT_MODEL).strip()
    if not MODEL_ID.match(model):
        raise ClaudeSdkError(f"模型：Bedrock 的模型或推理配置 ID，例如 {DEFAULT_MODEL}")
    fast = body.get("fastModel")
    fast = DEFAULT_FAST_MODEL if fast is None else str(fast).strip()
    if fast and not MODEL_ID.match(fast):
        raise ClaudeSdkError("后台模型：Bedrock 的模型或推理配置 ID（留空则用主模型）")
    tools_raw = body.get("tools") if body.get("tools") is not None else body.get("allowedTools")
    tools_raw = [] if tools_raw is None else tools_raw
    if not isinstance(tools_raw, (list, tuple)):
        raise ClaudeSdkError(f"内置工具：一个列表（{'、'.join(BUILTIN_TOOLS)}）")
    unknown = [str(t) for t in tools_raw if str(t) not in BUILTIN_TOOLS]
    if unknown:
        hint = "（Bedrock 上没有 WebSearch）" if "WebSearch" in unknown else ""
        raise ClaudeSdkError(f"内置工具 {'、'.join(unknown)} 不能用{hint}：可选 {'、'.join(BUILTIN_TOOLS)}")
    tools = [t for t in BUILTIN_TOOLS if t in {str(x) for x in tools_raw}]
    if not isinstance(body.get("skills") or [], (list, tuple)) or not isinstance(body.get("knowledgeBases") or [], (list, tuple)):
        raise ClaudeSdkError("技能和知识库：各是一个列表")
    skills: list[dict[str, Any]] = []
    for raw in body.get("skills") or []:
        ref = f"{raw.get('name')}{'@' + raw['version'] if raw.get('version') else ''}" if isinstance(raw, Mapping) else str(raw)
        found = SKILL_REF.match(ref.strip())
        if not found:
            raise ClaudeSdkError(f"技能 {ref!r}：技能库里的名称（小写字母、数字、连字符），可以带版本 @v0001")
        if found["name"] not in [s["name"] for s in skills]:
            skills.append({"name": found["name"], "version": found["version"]})
    if len(skills) > MAX_SKILLS:
        raise ClaudeSdkError(f"技能最多 {MAX_SKILLS} 个")
    kbs: list[str] = []
    for raw in body.get("knowledgeBases") or []:
        kb_id = str(raw.get("id") if isinstance(raw, Mapping) else raw).strip()
        if not deploy.KB_ID.match(kb_id):
            raise ClaudeSdkError(f"知识库 {kb_id!r}：知识库 ID（10 个字母或数字）")
        if kb_id not in kbs:
            kbs.append(kb_id)
    if len(kbs) > MAX_KBS:
        raise ClaudeSdkError(f"知识库最多 {MAX_KBS} 个")
    effort = str(body.get("effort") or "").strip().lower() or None
    if effort and effort not in EFFORTS:
        raise ClaudeSdkError(f"推理强度：{'、'.join(EFFORTS)}，或留空用模型的默认")
    budget = body.get("maxBudgetUsd")
    if budget in (None, ""):
        budget = None
    else:
        try:
            budget = round(float(budget), 4)
        except (TypeError, ValueError) as exc:
            raise ClaudeSdkError("每轮费用上限：美元数，例如 0.5（留空不限）") from exc
        if not 0.01 <= budget <= 100:
            raise ClaudeSdkError("每轮费用上限：0.01 到 100 美元")
    return {"name": name, "description": _line(body.get("description"), 300), "systemPrompt": prompt, "model": model, "fastModel": fast,
            "maxTurns": _int(body.get("maxTurns"), DEFAULT_TURNS, TURNS_MIN, TURNS_MAX, "最多轮数"), "tools": tools, "skills": skills,
            "knowledgeBases": kbs, "topK": _int(body.get("topK"), DEFAULT_TOP_K, TOP_K_MIN, TOP_K_MAX, "每次检索的段落数"), "effort": effort,
            "maxBudgetUsd": budget, "sample": _clean(body.get("sample"), MAX_SAMPLE).strip()}


def spec_hash(spec: Mapping[str, Any]) -> str:
    return hashlib.sha256(json.dumps(spec, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


# -- what the spec needs from AWS: its skills' text and its knowledge bases ---------------------------------------------

def load_skills(console: Any, workspace: str, spec: Mapping[str, Any]) -> dict[str, dict[str, Any]]:
    """Each skill's SKILL.md from the workspace's skill library (the current version unless the spec pins one)."""
    out: dict[str, dict[str, Any]] = {}
    for ref in spec["skills"]:
        try:
            got = skills_lab.get_skill(console, workspace, ref["name"], ref.get("version"))
        except skills_lab.SkillLabError as exc:
            raise ClaudeSdkError(f"技能 {ref['name']}：{exc}（先在「技能」页导入或保存它）") from exc
        meta, _body = skills_lab.frontmatter(got["text"])
        if meta.get("name") != ref["name"]:
            raise ClaudeSdkError(f"技能 {ref['name']} 的 SKILL.md 写的名字是 {meta.get('name')}：Claude Code 按目录名找它，两者要一致")
        out[ref["name"]] = {"name": ref["name"], "version": got["version"], "text": got["text"], "description": meta.get("description") or ""}
    return out


def resolve_kbs(session: Any, region: str, ids: Sequence[str]) -> dict[str, dict[str, Any]]:
    """The knowledge bases as AWS describes them: each must exist, be ACTIVE and take Retrieve (MANAGED or VECTOR)."""
    out: dict[str, dict[str, Any]] = {}
    if not ids:
        return out
    agent = client(session, "bedrock-agent", region)
    for kb_id in ids:
        try:
            kb = agent.get_knowledge_base(knowledgeBaseId=kb_id)["knowledgeBase"]
        except Exception as exc:  # noqa: BLE001
            if deploy._code(exc) == "ResourceNotFoundException":
                raise ClaudeSdkError(f"知识库 {kb_id} 不在这个工作区（{region}）") from exc
            raise
        kind = str((kb.get("knowledgeBaseConfiguration") or {}).get("type") or "")
        if kb.get("status") != "ACTIVE":
            raise ClaudeSdkError(f"知识库 {kb.get('name')}（{kb_id}）现在是 {kb.get('status')}：等它 ACTIVE 再部署")
        if kind not in ("MANAGED", "VECTOR"):
            raise ClaudeSdkError(f"知识库 {kb.get('name')}（{kb_id}）是 {kind} 类型：检索工具只支持 MANAGED 和 VECTOR")
        out[kb_id] = {"id": kb_id, "name": str(kb.get("name") or kb_id), "type": kind, "description": _line(kb.get("description"), 300)}
    return out


# -- generation -----------------------------------------------------------------------------------------------------------

def system_prompt(spec: Mapping[str, Any], kbs: Mapping[str, Mapping[str, Any]]) -> str:
    """The spec's prompt, then what the agent should know about its knowledge bases (skills list themselves)."""
    text = spec["systemPrompt"]
    if spec["knowledgeBases"]:
        lines = ["", "", "## 知识库", f"你有一个知识库检索工具 {KB_TOOL_NAME}（{KB_TOOL}）。问题涉及下面这些知识库的内容时，先检索，再按检索到的原文回答并说明出处；"
                 "检索不到就直说，不要编造。"]
        for kb_id in spec["knowledgeBases"]:
            kb = kbs.get(kb_id) or {"name": kb_id, "description": ""}
            lines.append(f"- {kb['name']}（kb_id {kb_id}）" + (f"：{kb['description']}" if kb.get("description") else ""))
        text += "\n".join(lines)
    return text


def kb_tool_description(spec: Mapping[str, Any], kbs: Mapping[str, Mapping[str, Any]]) -> str:
    names = "、".join(f"「{(kbs.get(k) or {}).get('name') or k}」" for k in spec["knowledgeBases"])
    more = "；kb_id 只检索其中一个（不填检索全部）" if len(spec["knowledgeBases"]) > 1 else ""
    return f"在知识库{names}里检索与问题相关的原文段落（Bedrock 知识库 Retrieve），返回段落、出处和相关度{more}。回答涉及知识库的内容前先调用。"


def dockerfile(spec: Mapping[str, Any], generated_at: str) -> str:
    return f"""# Generated by the ADLC console (Claude Agent SDK template): {spec['name']} · {generated_at}
# AgentCore Runtime runs linux/arm64 containers; the console's CodeBuild builds this on ARM and the deploy's scan
# gate reads the image's ECR scan before a runtime runs it.
FROM --platform=linux/arm64 {BASE_IMAGE}

# The base image's security updates, a font with Chinese for charts, Node {NODE_VERSION} from nodejs.org (checked
# against its SHA-256) and the Claude Code CLI that claude-agent-sdk drives, pinned to the version the SDK release bundles.
ARG CLAUDE_CODE_VERSION={CLAUDE_CODE}
ARG NODE_VERSION={NODE_VERSION}
ARG NODE_SHA256={NODE_SHA256}
RUN apt-get update \\
    && apt-get upgrade -y \\
    && apt-get install -y --no-install-recommends ca-certificates curl xz-utils {CHART_FONT_PACKAGE} \\
    && curl -fsSLo /tmp/node.tar.xz "https://nodejs.org/dist/v${{NODE_VERSION}}/node-v${{NODE_VERSION}}-linux-arm64.tar.xz" \\
    && echo "${{NODE_SHA256}}  /tmp/node.tar.xz" | sha256sum -c - \\
    && tar -xJf /tmp/node.tar.xz -C /usr/local --strip-components=1 --no-same-owner \\
    && rm /tmp/node.tar.xz \\
    && npm install --global --no-fund --no-audit "@anthropic-ai/claude-code@${{CLAUDE_CODE_VERSION}}" \\
    && npm cache clean --force \\
    && apt-get purge -y curl xz-utils \\
    && apt-get autoremove -y \\
    && rm -rf /var/lib/apt/lists/* /root/.npm \\
    && node --version && claude --version

WORKDIR /app
COPY requirements.txt .
# The SDK's wheel carries its own copy of the same CLI (225 MB): dropped, the agent runs the one on PATH.
# Retrieve on a managed knowledge base needs boto3 >= 1.43.90: checked here, not at the first question.
RUN pip install --no-cache-dir -r requirements.txt \\
    && python -c "import pathlib, claude_agent_sdk; (pathlib.Path(claude_agent_sdk.__file__).parent / '_bundled' / 'claude').unlink(missing_ok=True)" \\
    && python -c "import boto3, bedrock_agentcore.runtime, claude_agent_sdk; assert tuple(map(int, boto3.__version__.split('.')[:3])) >= (1, 43, 90), boto3.__version__" \\
    && python -c "import docx, matplotlib, openpyxl, PIL, pptx, pypdf"

COPY . .
# Claude Code refuses to bypass permissions as root, and an agent has no business being root: a user of its own.
RUN useradd --uid 1001 --home-dir /app --no-create-home --shell /bin/bash agent \\
    && mkdir -p /app/workspace \\
    && chown -R agent:agent /app
USER agent

# Bedrock with the runtime's execution role: no API key anywhere.
ENV CLAUDE_CODE_USE_BEDROCK=1 \\
    HOME=/app \\
    SHELL=/bin/bash \\
    PYTHONUNBUFFERED=1 \\
    CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC=1 \\
    DISABLE_AUTOUPDATER=1 \\
    MPLBACKEND=Agg \\
    MPLCONFIGDIR=/app/.matplotlib

# matplotlib draws with the font that has Chinese (and Latin) glyphs; its font cache is built here, not at a turn.
RUN mkdir -p "$MPLCONFIGDIR" \\
    && printf 'font.family: sans-serif\\nfont.sans-serif: {CHART_FONT}, DejaVu Sans\\naxes.unicode_minus: False\\n' > "$MPLCONFIGDIR/matplotlibrc" \\
    && python -c "import matplotlib.font_manager as fm, matplotlib.pyplot; assert '{CHART_FONT}' in {{f.name for f in fm.fontManager.ttflist}}"

EXPOSE 8080
CMD ["python", "main.py"]
"""


def requirements() -> str:
    return "\n".join(["# generated by the ADLC console (Claude Agent SDK template): installed by the image's own pip, linux/arm64",
                      SDK, AGENTCORE, *BOTO, *FILE_READERS, *CHARTS]) + "\n"


MAIN_TEMPLATE = r'''__DOC__
import asyncio
import base64
import csv
import hashlib
import io
import json
import logging
import os
import re
import secrets
import shutil
import stat
import threading
import time
import uuid
import zipfile
from pathlib import Path

import boto3
from bedrock_agentcore.runtime import BedrockAgentCoreApp
from claude_agent_sdk import (AssistantMessage, ClaudeAgentOptions, ResultMessage, TextBlock, ToolResultBlock, ToolUseBlock, UserMessage,
                              create_sdk_mcp_server, query, tool)
from starlette.responses import Response

AGENT = __AGENT__
MODEL = __MODEL__
#: Claude Code's background model (titles, WebFetch's extraction): ANTHROPIC_DEFAULT_HAIKU_MODEL ("" = the main model).
FAST_MODEL = __FAST_MODEL__
MAX_TURNS = __MAX_TURNS__
EFFORT = __EFFORT__
MAX_BUDGET_USD = __MAX_BUDGET_USD__
#: The built-in tools this agent has: nothing else is offered to the model (``tools``), and these run unasked.
BUILTIN_TOOLS = __BUILTIN_TOOLS__
#: Skills from the console's library, in .claude/skills/<name>/SKILL.md: the model loads one through the Skill tool.
SKILLS = __SKILLS__
KNOWLEDGE_BASES = __KNOWLEDGE_BASES__
KB_TOP_K = __KB_TOP_K__
KB_DESCRIPTION = __KB_DESCRIPTION__
SYSTEM_PROMPT = __SYSTEM_PROMPT__
REGION = os.environ.get("AWS_REGION") or os.environ.get("AWS_DEFAULT_REGION") or __REGION__
KB_SERVER = "kb"
KB_TOOL = f"mcp__{KB_SERVER}__retrieve"
APP_DIR = Path(__file__).resolve().parent
SKILLS_SOURCE = APP_DIR / ".claude" / "skills"
#: A runtime session's Claude Code session id is uuid5(this, the runtime session id) on its first turn.
SESSION_NAMESPACE = uuid.UUID("6f1c0a52-3d0e-4b8f-9a77-0c5d2b1e9a41")
MAX_PROMPT = 20000

#: A call with its own skills (an evaluation's: the console's Skill Lab): ``skills`` names what it may load, each
#: ``skillOverrides`` entry is a skill version's folder in S3, read for this call only.
SKILL_NAME = re.compile(r"^[a-z0-9](?:[a-z0-9-]{0,62}[a-z0-9])?$")
SKILL_URI = re.compile(r"^s3://([a-z0-9][a-z0-9.-]{1,61}[a-z0-9])/((?:[^/]+/)+)$")
MODEL_ID = re.compile(r"^(?:[a-z0-9][a-z0-9.:_\-/]{2,199}|arn:aws[a-z-]*:bedrock:[a-z0-9-]+:\d{12}:[a-z-]+/[A-Za-z0-9._:/-]{1,200})$")
MAX_CALL_SKILLS, MAX_SKILL_FILES, MAX_SKILL_BYTES = 20, 200, 10_000_000
#: A call's input files (``files``: an evaluation task's): each read from S3 into the call's working directory before
#: the prompt (its size and SHA-256 checked), at most MAX_INPUTS of at most MAX_INPUT_FILE, MAX_INPUT_BYTES in all.
#: They are in the directory's manifest before the turn: one the turn leaves alone is not one of its files.
INPUT_URI = re.compile(r"^s3://([a-z0-9][a-z0-9.-]{1,61}[a-z0-9])/([^?#\s]{1,1024})$")
SHA256 = re.compile(r"^[0-9a-f]{64}$")
CONTROL = re.compile(r"[\x00-\x1f\x7f]")
MAX_INPUTS, MAX_INPUT_FILE, MAX_INPUT_BYTES = 32, 25 * 1024 * 1024, 100 * 1024 * 1024

#: The files a turn created or changed, in its answer: at most MAX_ARTIFACTS, each with a text view of MAX_TEXT
#: characters (MAX_TEXT_TOTAL in all) and, up to MAX_INLINE bytes (MAX_INLINE_TOTAL in all), its bytes in base64.
MAX_ARTIFACTS, MAX_TEXT, MAX_TEXT_TOTAL, MAX_INLINE, MAX_INLINE_TOTAL = 20, 6000, 24000, 48 * 1024, 160 * 1024
MAX_SCAN, MAX_PARSE = 5000, 20 * 1024 * 1024
#: A spreadsheet's (or a CSV's, a Word table's) preview: the first sheets, rows and columns, each cell shortened.
TABLE_SHEETS, TABLE_ROWS, TABLE_COLS, CELL_CHARS = 5, 40, 15, 120
#: A presentation's view: its first slides.
MAX_SLIDES = 40
#: Pictures for whoever judges the turn: an image the turn made, and the pictures inside a document it made (at most
#: DOC_PICTURES a document), each as a preview of at most PREVIEW_PX on its long side and PREVIEW_BYTES (a model looks
#: at about 1568 px) and a thumbnail for the console's page; at most MAX_PICTURES previews, PICTURES_BYTES in all.
PREVIEW_PX, PREVIEW_BYTES, THUMB_PX, THUMB_BYTES = 1568, 1_000_000, 320, 64 * 1024
MAX_PICTURES, PICTURES_BYTES, DOC_PICTURES, MAX_PIXELS = 8, 4_000_000, 4, 50_000_000
CONVERSE_FORMATS = ("png", "jpeg", "gif", "webp")
#: Where a Word document or a workbook keeps its pictures (a presentation's come from its slides, in order).
MEDIA = {"docx": "word/media/", "xlsx": "xl/media/"}
#: Directories that are Claude Code's or a tool's own, never the agent's work.
SKIP_DIRS = {".claude-agent", ".claude", ".git", "__pycache__", "node_modules", ".cache", ".npm", ".matplotlib"}
KINDS = {".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv", ".tsv": "csv", ".md": "md", ".markdown": "md", ".txt": "txt", ".log": "txt",
         ".json": "json", ".jsonl": "json", ".pdf": "pdf", ".docx": "docx", ".html": "html", ".htm": "html", ".xml": "xml", ".yaml": "yaml",
         ".yml": "yaml", ".py": "code", ".js": "code", ".ts": "code", ".sh": "code", ".sql": "code", ".png": "image", ".jpg": "image",
         ".jpeg": "image", ".gif": "image", ".webp": "image", ".bmp": "image", ".tif": "image", ".tiff": "image", ".svg": "xml", ".pptx": "pptx",
         ".xls": "xls", ".doc": "doc", ".zip": "zip"}
TEXT_KINDS = {"md", "txt", "json", "html", "xml", "yaml", "code"}
NUMBERISH = re.compile(r"^[+-]?\d[\d,]*(?:\.\d+)?$")

logger = logging.getLogger("claude_agent")  # its own handler: bedrock_agentcore logs through its own
logger.setLevel(logging.INFO)
logger.propagate = False
if not logger.handlers:
    logger.addHandler(logging.StreamHandler())
    logger.handlers[0].setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))

app = BedrockAgentCoreApp()
_SETUP = {}
_SETUP_LOCK = threading.Lock()
_LOCKS = {}


class PayloadError(ValueError):
    """A payload the caller must fix: answered 400."""


class SkillFetchError(RuntimeError):
    """A skill override that could not be read from S3: answered 503 (a read grant made seconds ago may not be visible yet)."""


# -- the workspace: session storage when the runtime has it ------------------------------------------------------------
def _writable(path):
    try:
        path.mkdir(parents=True, exist_ok=True)
        probe = path / ".write-check"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink()
        return True
    except OSError as exc:
        logger.warning("%s is not writable: %s", path, exc)
        return False


def _setup():
    """Once per process: the workspace (AGENT_WORKSPACE, the runtime's session storage, when it is writable), Claude
    Code's config dir in it (its session transcripts survive a stopped session there) and the image's skills."""
    with _SETUP_LOCK:
        if _SETUP:
            return _SETUP
        wanted = os.environ.get("AGENT_WORKSPACE") or ""
        persistent = bool(wanted) and _writable(Path(wanted))
        workspace = Path(wanted) if persistent else APP_DIR / "workspace"
        workspace.mkdir(parents=True, exist_ok=True)
        config = workspace / ".claude-agent"
        config.mkdir(parents=True, exist_ok=True)
        target = config / "skills"
        if target.exists():
            shutil.rmtree(target)
        if SKILLS_SOURCE.is_dir():
            shutil.copytree(SKILLS_SOURCE, target)
        _SETUP.update(workspace=workspace, persistent=persistent, config=config, cli=shutil.which("claude"))
        logger.info("ready %s", json.dumps({"agent": AGENT.get("name"), "workspace": str(workspace), "sessionStorage": persistent,
                                            "skills": sorted(p.name for p in target.iterdir()) if target.is_dir() else [],
                                            "cli": _SETUP["cli"], "model": MODEL, "format": AGENT.get("format")}, ensure_ascii=False))
        return _SETUP


def _conversations(setup):
    path = setup["config"] / "conversations.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _remember(setup, runtime_session, claude_session):
    found = _conversations(setup)
    found[runtime_session] = claude_session
    path = setup["config"] / "conversations.json"
    path.write_text(json.dumps(dict(list(found.items())[-200:])), encoding="utf-8")


def _has_transcript(setup, claude_session):
    return any((setup["config"] / "projects").glob(f"*/{claude_session}.jsonl"))


# -- a call's own skills and input files -------------------------------------------------------------------------------------
def _call_plan(payload):
    """What a call brings, when it brings anything: ``None`` (the image's skills, in the session's workspace: a
    conversation) or ``{"names", "overrides", "files"}``, a call of its own. ``skills``: the names this call may load
    ([] = none; left out = the image's); ``skillOverrides``: ``[{"name", "uri": "s3://bucket/prefix/"}]``, each read
    from S3 for this call; ``files``: its input files (see :func:`_call_files`)."""
    names, raw, files = payload.get("skills"), payload.get("skillOverrides"), payload.get("files")
    if names is None and not raw and not files:
        return None
    if names is not None and (not isinstance(names, list) or not all(isinstance(n, str) for n in names)):
        raise PayloadError("skills: a list of skill names ([] for none)")
    if raw is not None and not isinstance(raw, list):
        raise PayloadError("skillOverrides: a list of {\"name\", \"uri\"}")
    overrides = {}
    for item in raw or []:
        name, uri = (str(item.get("name") or ""), str(item.get("uri") or "")) if isinstance(item, dict) else ("", "")
        if not SKILL_NAME.match(name) or not SKILL_URI.match(uri):
            raise PayloadError("skillOverrides: each one {\"name\": a skill name, \"uri\": \"s3://bucket/prefix/\"} (SKILL.md at the prefix)")
        overrides[name] = uri
    baked = [s["name"] for s in SKILLS]
    chosen = list(dict.fromkeys(names if names is not None else baked + list(overrides)))
    unknown = [n for n in chosen if n not in overrides and n not in baked]
    if unknown:
        raise PayloadError(f"skills: {', '.join(unknown)} is not one of this agent's skills ({', '.join(baked) or 'none'}) and has no override")
    stray = [n for n in overrides if n not in chosen]
    if stray:
        raise PayloadError(f"skillOverrides: {', '.join(stray)} is not one of this call's skills")
    if len(chosen) > MAX_CALL_SKILLS:
        raise PayloadError(f"skills: at most {MAX_CALL_SKILLS}")
    return {"names": chosen, "overrides": overrides, "files": _call_files(files)}


def _call_files(raw):
    """A call's input files: ``[{"path", "uri": "s3://bucket/key"[, "sha256", "bytes"]}]``, each ``path`` a plain
    relative path in the call's directory (at most 4 levels, no hidden names)."""
    if raw is None:
        return []
    if not isinstance(raw, list):
        raise PayloadError('files: a list of {"path", "uri"}')
    if len(raw) > MAX_INPUTS:
        raise PayloadError(f"files: at most {MAX_INPUTS}")
    out, seen, declared = [], set(), 0
    for item in raw:
        if not isinstance(item, dict):
            raise PayloadError('files: each one {"path", "uri"[, "sha256", "bytes"]}')
        path, uri = str(item.get("path") or ""), str(item.get("uri") or "")
        parts = path.split("/")
        if (not path or len(path) > 240 or len(parts) > 4 or path.startswith(("/", "~")) or "\\" in path or CONTROL.search(path)
                or any(p in ("", ".", "..") or p.startswith(".") for p in parts)):
            raise PayloadError(f"files: {path!r} is not a plain relative path (at most 4 levels, no '..', no hidden names)")
        if path.casefold() in seen:
            raise PayloadError(f"files: {path} is named twice")
        seen.add(path.casefold())
        if not INPUT_URI.match(uri) or any(p in (".", "..") for p in uri[5:].split("/")):
            raise PayloadError(f"files: {path}: the uri must be s3://bucket/key")
        sha = str(item.get("sha256") or "").lower()
        if sha and not SHA256.match(sha):
            raise PayloadError(f"files: {path}: sha256 is 64 hex digits")
        size = item.get("bytes")
        if size is not None and (isinstance(size, bool) or not isinstance(size, int) or not 0 <= size <= MAX_INPUT_FILE):
            raise PayloadError(f"files: {path}: bytes is a size of at most {MAX_INPUT_FILE // 1024 // 1024} MiB")
        declared += size or 0
        if declared > MAX_INPUT_BYTES:
            raise PayloadError(f"files: at most {MAX_INPUT_BYTES // 1024 // 1024} MiB in all")
        out.append({"path": path, "uri": uri, "sha256": sha or None, "bytes": size})
    return out


def _fetch_inputs(files, workdir):
    """Each input file from S3 to its path in the call's directory, its size and SHA-256 checked against the call's."""
    s3, placed, total = _s3(), [], 0
    for item in files:
        bucket, key = INPUT_URI.match(item["uri"]).groups()
        target = workdir.joinpath(*item["path"].split("/"))
        target.parent.mkdir(parents=True, exist_ok=True)
        digest, size = hashlib.sha256(), 0
        try:
            body = s3.get_object(Bucket=bucket, Key=key)["Body"]
            try:
                with target.open("wb") as out:
                    while True:
                        chunk = body.read(1024 * 1024)
                        if not chunk:
                            break
                        size += len(chunk)
                        total += len(chunk)
                        if size > MAX_INPUT_FILE or total > MAX_INPUT_BYTES:
                            raise PayloadError(f"files: {item['path']} is larger than the call says it is (or than {MAX_INPUT_FILE // 1024 // 1024} MiB)")
                        digest.update(chunk)
                        out.write(chunk)
            finally:
                if hasattr(body, "close"):
                    body.close()
        except PayloadError:
            raise
        except Exception as exc:  # noqa: BLE001 - AccessDenied (a grant not visible yet), a throttle, the network: try again
            code = str(((getattr(exc, "response", None) or {}).get("Error") or {}).get("Code") or "")
            if code in ("NoSuchKey", "404") or "NoSuchKey" in str(exc):
                raise PayloadError(f"files: {item['uri']} does not exist") from exc
            raise SkillFetchError(f"files: {item['uri']} could not be read: {type(exc).__name__}: {str(exc)[:300]}") from exc
        sha = digest.hexdigest()
        if (item["sha256"] and sha != item["sha256"]) or (item["bytes"] is not None and size != item["bytes"]):
            raise PayloadError(f"files: {item['uri']} is not the file the call names for {item['path']} (sha256 {sha[:12]}…, {size} bytes)")
        placed.append({"path": item["path"], "bytes": size, "sha256": sha})
    return placed


def _inputs_note(inputs):
    """What the model is told of a call's input files, after the prompt."""
    listed = "、".join(f"{i['path']}（{_size(i['bytes'])}）" for i in inputs)
    return f"〔输入文件〕这次任务的文件已经放在当前工作目录里：{listed}。"


def _call_model(payload):
    model = str(payload.get("model") or "").strip()
    if model and not MODEL_ID.match(model):
        raise PayloadError("model: a Bedrock model or inference profile id")
    return model or None


_S3 = None
_S3_LOCK = threading.Lock()


def _s3():
    global _S3
    with _S3_LOCK:
        if _S3 is None:
            _S3 = boto3.client("s3", region_name=REGION)
        return _S3


def _frontmatter_name(text):
    head = re.match(r"^\ufeff?---[ \t]*\r?\n(.*?)\r?\n---", text, re.S)
    found = re.search(r"^name:[ \t]*(.+?)[ \t]*$", head.group(1), re.M) if head else None
    return found.group(1).strip("\"'") if found else None


def _fetch_skill(name, uri, target):
    """The skill folder under ``uri`` (its SKILL.md at the prefix's root, naming the skill) into ``target``."""
    bucket, prefix = SKILL_URI.match(uri).groups()
    s3, token, keys, total = _s3(), None, [], 0
    try:
        while len(keys) <= MAX_SKILL_FILES:
            page = s3.list_objects_v2(Bucket=bucket, Prefix=prefix, **({"ContinuationToken": token} if token else {}))
            for found in page.get("Contents") or []:
                rel = str(found["Key"])[len(prefix):]
                if rel and not rel.endswith("/") and not any(part in ("", ".", "..") for part in rel.split("/")):
                    keys.append((found["Key"], rel))
                    total += int(found.get("Size") or 0)
            token = page.get("NextContinuationToken")
            if not page.get("IsTruncated") or not token:
                break
        if len(keys) > MAX_SKILL_FILES or total > MAX_SKILL_BYTES:
            raise PayloadError(f"skillOverrides: {uri} holds more than {MAX_SKILL_FILES} files or {MAX_SKILL_BYTES // 1_000_000} MB")
        if not any(rel == "SKILL.md" for _key, rel in keys):
            raise PayloadError(f"skillOverrides: no SKILL.md at {uri}")
        for key, rel in keys:
            path = target.joinpath(*rel.split("/"))
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(s3.get_object(Bucket=bucket, Key=key)["Body"].read())
    except PayloadError:
        raise
    except Exception as exc:  # noqa: BLE001 - AccessDenied (a grant not visible yet), a throttle, the network
        raise SkillFetchError(f"skillOverrides: {uri} could not be read: {type(exc).__name__}: {str(exc)[:300]}") from exc
    named = _frontmatter_name((target / "SKILL.md").read_text(encoding="utf-8", errors="replace"))
    if named != name:
        raise PayloadError(f"skillOverrides: the SKILL.md at {uri} names {named or 'no skill'}, not {name}")
    return len(keys)


def _prepare_call(setup, plan):
    """A call of its own: its own working directory (``calls/<id>`` in the workspace, so what it writes is its own),
    holding its input files, and its own Claude Code config whose skills are exactly the call's (the image's copied
    in, each override read from S3). Claude Code reads that config's settings only (``setting_sources=["user"]``): no
    project skill of the workspace or of the image joins in."""
    call_id = time.strftime("%Y%m%d-%H%M%S", time.gmtime()) + "-" + secrets.token_hex(4)
    workdir = setup["workspace"] / "calls" / call_id
    config = setup["config"] / "calls" / call_id
    (config / "skills").mkdir(parents=True)
    workdir.mkdir(parents=True)
    versions = {s["name"]: s.get("version") for s in SKILLS}
    used = []
    for name in plan["names"]:
        if name in plan["overrides"]:
            files = _fetch_skill(name, plan["overrides"][name], config / "skills" / name)
            used.append({"name": name, "source": plan["overrides"][name], "files": files})
        else:
            shutil.copytree(SKILLS_SOURCE / name, config / "skills" / name)
            used.append({"name": name, "source": "image", "version": versions.get(name)})
    inputs = _fetch_inputs(plan.get("files") or [], workdir)
    return {"id": call_id, "workdir": workdir, "config": config, "skills": used, "inputs": inputs}


# -- the files a turn created or changed ------------------------------------------------------------------------------------
def _manifest(root, skip=()):
    """``{path: (size, mtime_ns)}`` of the regular files under ``root`` (links, SKIP_DIRS and the top-level ``skip``
    left out), at most MAX_SCAN."""
    out = {}
    for folder, dirs, names in os.walk(root):
        here = Path(folder)
        dirs[:] = sorted(d for d in dirs if d not in SKIP_DIRS and not (here == root and d in skip))
        for name in sorted(names):
            path = here / name
            try:
                info = path.lstat()
            except OSError:
                continue
            if stat.S_ISREG(info.st_mode):
                out[path.relative_to(root).as_posix()] = (info.st_size, info.st_mtime_ns)
                if len(out) >= MAX_SCAN:
                    return out
    return out


def _kind(path):
    kind = KINDS.get(path.suffix.lower())
    if kind:
        return kind
    try:
        with path.open("rb") as handle:
            head = handle.read(2048)
    except OSError:
        return "binary"
    if b"\x00" in head:
        return "binary"
    try:
        head.decode("utf-8")
    except UnicodeDecodeError as exc:
        if exc.start < len(head) - 4:  # not just a character cut off at the end of the sample
            return "binary"
    return "txt"


def _read_text(path, limit=MAX_TEXT + 1):
    with path.open("rb") as handle:
        raw = handle.read(limit * 4)
    return raw.decode("utf-8", "replace").lstrip("\ufeff")[:limit]


def _cell(value, cached=None):
    """A spreadsheet cell as text: a formula as written (with its last computed value when the file has one), a
    string that looks like a number in quotes (it is text in the file), dates in ISO form."""
    if value is None:
        return ""
    formula = getattr(value, "text", None)  # an array formula
    if isinstance(formula, str):
        value = formula
    if isinstance(value, str) and value.startswith("="):
        got = _cell(cached) if cached is not None and not (isinstance(cached, str) and cached.startswith("=")) else ""
        text = value + (f" → {got}" if got else "")
    elif isinstance(value, str):
        text = json.dumps(value, ensure_ascii=False) if NUMBERISH.match(value.strip()) else value
    elif isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif hasattr(value, "isoformat"):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.replace("\r\n", "\n").replace("\n", "\\n")
    return text if len(text) <= CELL_CHARS else text[:CELL_CHARS] + "…"


def _col(number):
    out = ""
    while number:
        number, rest = divmod(number - 1, 26)
        out = chr(65 + rest) + out
    return out


def _grid(rows):
    if not rows:
        return ["（空）"]
    width = max(len(r) for r in rows)
    return ["    | " + " | ".join(_col(i + 1) for i in range(width))] + [f"{i:>3} | " + " | ".join(row) for i, row in enumerate(rows, 1)]


def _xlsx(path):
    import openpyxl  # in the image (requirements.txt): a turn's spreadsheet, read for whoever judges it

    book = openpyxl.load_workbook(path, data_only=False)
    try:
        values = openpyxl.load_workbook(path, data_only=True)  # a formula's last computed value, when the file holds one
    except Exception:  # noqa: BLE001
        values = None
    lines, sheets = [], []
    for index, sheet in enumerate(book.worksheets):
        rows_n, cols_n = int(sheet.max_row or 0), int(sheet.max_column or 0)
        if rows_n <= 1 and cols_n <= 1 and sheet.cell(1, 1).value is None:
            rows_n = cols_n = 0
        cached = values[sheet.title] if values is not None and sheet.title in values.sheetnames else None
        shown_rows, shown_cols = min(rows_n, 60), min(cols_n, 20)
        grid = [[_cell(sheet.cell(r, c).value, cached.cell(r, c).value if cached is not None else None) for c in range(1, shown_cols + 1)]
                for r in range(1, shown_rows + 1)]
        notes = [f"{rows_n} 行 × {cols_n} 列"]
        if sheet.freeze_panes:
            notes.append(f"冻结窗格 {sheet.freeze_panes}")
        header = [sheet.cell(1, c) for c in range(1, shown_cols + 1) if sheet.cell(1, c).value is not None]
        if header and all(cell.font is not None and cell.font.b for cell in header):
            notes.append("首行加粗")
        merged = len(sheet.merged_cells.ranges)
        if merged:
            notes.append(f"合并单元格 {merged} 处")
        lines.append(f"[工作表 {index + 1}：{sheet.title}] " + " · ".join(notes))
        lines += _grid(grid)
        if rows_n > shown_rows or cols_n > shown_cols:
            lines.append(f"（只列出前 {shown_rows} 行、{shown_cols} 列）")
        lines.append("")
        if index < TABLE_SHEETS:
            sheets.append({"name": sheet.title, "rows": [r[:TABLE_COLS] for r in grid[:TABLE_ROWS]], "nrows": rows_n, "ncols": cols_n,
                           "truncated": rows_n > TABLE_ROWS or cols_n > TABLE_COLS})
    return {"text": "\n".join(lines).strip(), "table": {"sheets": sheets}, "sheets": len(book.worksheets)}


def _csv(path):
    with path.open("rb") as handle:
        text = handle.read(2_000_000).decode("utf-8-sig", "replace")
    first = text.split("\n", 1)[0]
    delimiter = "\t" if path.suffix.lower() == ".tsv" or ("\t" in first and "," not in first) else ";" if ";" in first and "," not in first else ","
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows.append(row)
        if len(rows) > 100_000:
            break
    width = max((len(r) for r in rows), default=0)
    shown = [[str(v).replace("\n", "\\n")[:CELL_CHARS] for v in r[:TABLE_COLS]] for r in rows[:TABLE_ROWS]]
    return {"text": text, "table": {"sheets": [{"name": path.name, "rows": shown, "nrows": len(rows), "ncols": width,
                                                "truncated": len(rows) > TABLE_ROWS or width > TABLE_COLS}]}}


def _docx(path):
    import docx  # python-docx

    document = docx.Document(str(path))
    lines, tables = [], []
    blocks = document.iter_inner_content() if hasattr(document, "iter_inner_content") else [*document.paragraphs, *document.tables]
    for block in blocks:
        if hasattr(block, "rows"):  # a table, in the document's order
            rows = list(block.rows)
            grid = [[cell.text.strip().replace("\n", " / ")[:CELL_CHARS] for cell in row.cells] for row in rows[:60]]
            tables.append({"name": f"表格 {len(tables) + 1}", "rows": [r[:TABLE_COLS] for r in grid[:TABLE_ROWS]], "nrows": len(rows),
                           "ncols": max((len(r) for r in grid), default=0), "truncated": len(rows) > TABLE_ROWS})
            lines.append(f"[表格 {len(tables)}]")
            lines += _grid(grid)
            continue
        text = str(getattr(block, "text", "") or "").strip()
        if not text:
            continue
        style = str(getattr(getattr(block, "style", None), "name", "") or "")
        level = re.match(r"^Heading (\d)", style)
        lines.append(("#" * int(level.group(1)) + " " if level else "# " if style == "Title" else "- " if "List" in style else "") + text)
    out = {"text": "\n".join(lines)}
    if tables:
        out["table"] = {"sheets": tables[:TABLE_SHEETS]}
    return out


def _pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        reader.decrypt("")
    count = len(reader.pages)
    parts = [f"[第 {i + 1} 页]\n{(reader.pages[i].extract_text() or '').strip()}" for i in range(min(count, 30))]
    return {"text": "\n\n".join(parts), "pages": count}


def _shapes(shapes, kinds):
    for shape in shapes:  # a group's shapes, in order
        if shape.shape_type == kinds.GROUP:
            yield from _shapes(shape.shapes, kinds)
        else:
            yield shape


def _chart(chart):
    kind = str(chart.chart_type).split(" (")[0]
    title = chart.chart_title.text_frame.text.strip() if chart.has_title and chart.chart_title.has_text_frame else ""
    parts = [f"[图表 {kind}]" + (f" 标题「{title}」" if title else "")]
    for plot in list(chart.plots)[:3]:
        categories = [str(c) for c in list(plot.categories)[:20]]
        if categories:
            parts.append("类别 " + "、".join(categories))
        for series in list(plot.series)[:8]:
            values = ["" if v is None else f"{v:g}" for v in list(series.values)[:20]]
            parts.append(f"系列「{series.name}」{', '.join(values)}")
    return "；".join(parts)


def _pptx(path):
    """A presentation's slides: the title, each text box's paragraphs (indented by level), tables as cells, a chart's
    type, title, categories and series, each picture (and its bytes, for a preview), the notes."""
    from pptx import Presentation  # python-pptx, in the image (requirements.txt)
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    slides = list(deck.slides)
    lines, tables, pictures = [], [], []
    for number, slide in enumerate(slides[:MAX_SLIDES], 1):
        title = slide.shapes.title
        heading = title.text_frame.text.strip() if title is not None and title.has_text_frame else ""
        lines.append(f"[幻灯片 {number}]" + (f" {heading}" if heading else ""))
        for shape in _shapes(slide.shapes, MSO_SHAPE_TYPE):
            if title is not None and shape.shape_id == title.shape_id:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip().replace("\x0b", " / ")
                    if text:
                        lines.append("  " * (int(paragraph.level or 0) + 1) + text)
            if shape.has_table:
                rows = list(shape.table.rows)
                grid = [[cell.text.strip().replace("\n", " / ")[:CELL_CHARS] for cell in row.cells] for row in rows[:60]]
                tables.append({"name": f"幻灯片 {number} · 表格 {len(tables) + 1}", "rows": [r[:TABLE_COLS] for r in grid[:TABLE_ROWS]],
                               "nrows": len(rows), "ncols": max((len(r) for r in grid), default=0), "truncated": len(rows) > TABLE_ROWS})
                lines.append(f"  [表格 {len(tables)}]")
                lines += ["  " + line for line in _grid(grid)]
            if shape.has_chart:
                lines.append("  " + _chart(shape.chart))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                try:
                    size = "{}×{} px".format(*image.size)
                except Exception:  # noqa: BLE001 - a picture Pillow cannot measure (an EMF)
                    size = "?"
                lines.append(f"  [图片 {shape.name} · {image.content_type} · {size}]")
                pictures.append((f"幻灯片 {number} · {shape.name}", image.blob))
        notes = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if notes is not None and notes.text.strip():
            lines.append(f"  备注：{notes.text.strip()}")
    if len(slides) > MAX_SLIDES:
        lines.append(f"（只列出前 {MAX_SLIDES} 张，共 {len(slides)} 张）")
    out = {"text": "\n".join(lines), "slides": len(slides)}
    if tables:
        out["table"] = {"sheets": tables[:TABLE_SHEETS]}
    if pictures:
        out.update(_pictures=pictures[:DOC_PICTURES], _pictureCount=len(pictures))
    return out


def _natural(name):
    return re.sub(r"\d+", lambda m: m.group(0).zfill(8), name)


def _media(path, kind):
    """The pictures a Word document or a workbook keeps (its media folder), the first DOC_PICTURES of them read."""
    with zipfile.ZipFile(path) as archive:
        infos = sorted((i for i in archive.infolist() if i.filename.startswith(MEDIA[kind]) and not i.is_dir()), key=lambda i: _natural(i.filename))
        pictures = [(i.filename, archive.read(i)) for i in infos[:DOC_PICTURES] if i.file_size <= MAX_PARSE]
    return {"_pictures": pictures, "_pictureCount": len(infos)} if infos else {}


def _image_format(data):
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    return None


def _flat(image):
    """The image on white, without alpha (for JPEG)."""
    from PIL import Image

    if image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        canvas = Image.new("RGB", rgba.size, (255, 255, 255))
        canvas.paste(rgba, mask=rgba.getchannel("A"))
        return canvas
    return image.convert("RGB")


def _encode(image, px, limit, photo=False):
    """The image at most ``px`` on its long side as PNG (a chart's lines stay sharp) or, over ``limit`` bytes — or
    first, for a ``photo`` — JPEG; halved until it fits (None when it does not)."""
    work = image.copy()
    work.thumbnail((px, px))
    if work.mode not in ("RGB", "RGBA", "L", "LA", "P"):
        work = work.convert("RGBA")
    for _ in range(4):
        if not photo:
            buffer = io.BytesIO()
            work.save(buffer, format="PNG", optimize=True)
            if buffer.tell() <= limit:
                return {"format": "png", "width": work.size[0], "height": work.size[1], "bytes": buffer.tell(),
                        "base64": base64.b64encode(buffer.getvalue()).decode("ascii")}
        flat = _flat(work)
        for quality in (85, 70):
            buffer = io.BytesIO()
            flat.save(buffer, format="JPEG", quality=quality, optimize=True)
            if buffer.tell() <= limit:
                return {"format": "jpeg", "width": work.size[0], "height": work.size[1], "bytes": buffer.tell(),
                        "base64": base64.b64encode(buffer.getvalue()).decode("ascii")}
        work = work.resize((max(1, work.size[0] // 2), max(1, work.size[1] // 2)))
    return None


def _picture(data, budget):
    """An image as the judge sees it and the console's page shows it: its format and size, a ``preview`` (at most
    PREVIEW_PX on its long side and PREVIEW_BYTES: the image itself when it already fits) and a ``thumb``, within
    the answer's ``budget`` of previews."""
    out = {"format": _image_format(data) or "unknown"}
    if budget["count"] <= 0 or budget["bytes"] <= 0:
        out["note"] = f"an answer previews at most {MAX_PICTURES} pictures"
        return out
    try:
        from PIL import Image, ImageOps  # Pillow, in the image (python-pptx's and matplotlib's too)
    except ImportError:  # an image without it: a small picture as it is
        if out["format"] in CONVERSE_FORMATS and len(data) <= min(PREVIEW_BYTES, budget["bytes"]):
            out["preview"] = {"format": out["format"], "bytes": len(data), "base64": base64.b64encode(data).decode("ascii")}
            budget.update(count=budget["count"] - 1, bytes=budget["bytes"] - len(data))
        return out
    Image.MAX_IMAGE_PIXELS = MAX_PIXELS
    try:
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            frames = int(getattr(image, "n_frames", 1) or 1)
            out.update(format=str(image.format or out["format"]).lower(), width=image.size[0], height=image.size[1])
            if frames > 1:
                out["frames"] = frames
            rotated = (image.getexif() or {}).get(0x0112, 1) not in (None, 1)
            frame = ImageOps.exif_transpose(image) if rotated else image
            if (out["format"] in CONVERSE_FORMATS and frames == 1 and not rotated and max(image.size) <= PREVIEW_PX
                    and len(data) <= min(PREVIEW_BYTES, budget["bytes"])):
                preview = {"format": out["format"], "width": image.size[0], "height": image.size[1], "bytes": len(data),
                           "base64": base64.b64encode(data).decode("ascii")}
            else:
                preview = _encode(frame, PREVIEW_PX, min(PREVIEW_BYTES, budget["bytes"]), photo=out["format"] in ("jpeg", "mpo"))
            thumb = _encode(frame, THUMB_PX, THUMB_BYTES, photo=out["format"] in ("jpeg", "mpo"))
    except Exception as exc:  # noqa: BLE001 - not a picture Pillow reads (an EMF in a document, a broken file): said so
        out["error"] = f"{type(exc).__name__}: {str(exc)[:200]}"
        return out
    if preview:
        out["preview"] = preview
        budget.update(count=budget["count"] - 1, bytes=budget["bytes"] - preview["bytes"])
    if thumb:
        out["thumb"] = thumb
    return out


def _view(path, kind, size):
    if kind in TEXT_KINDS:
        return {"text": _read_text(path)}
    if kind == "csv":
        return _csv(path)
    if kind in ("xlsx", "docx", "pdf", "pptx"):
        if size > MAX_PARSE:
            return {"error": f"larger than {MAX_PARSE // 1024 // 1024} MB: not read"}
        view = {"xlsx": _xlsx, "docx": _docx, "pdf": _pdf, "pptx": _pptx}[kind](path)
        if kind in ("xlsx", "docx"):
            try:
                view.update(_media(path, kind))
            except (OSError, zipfile.BadZipFile, KeyError, ValueError):
                pass
        return view
    return {}


def _artifacts(root, before, skip=()):
    """The files under ``root`` the turn created or changed (against ``before``, its manifest when the turn began):
    path, bytes, type, created/changed, a text view (a spreadsheet's sheets as cells, formulas as written; a PDF's,
    a Word document's or a presentation's text; a text file as it is) and a table preview, an image's (or a
    document's pictures') preview and thumbnail, the bytes in base64 when small."""
    after = _manifest(root, skip)
    paths = sorted(p for p, sig in after.items() if before.get(p) != sig)
    out, text_left, inline_left = [], MAX_TEXT_TOTAL, MAX_INLINE_TOTAL
    budget = {"count": MAX_PICTURES, "bytes": PICTURES_BYTES}
    for rel in paths[:MAX_ARTIFACTS]:
        path, size = root / rel, after[rel][0]
        kind = _kind(path)
        item = {"path": rel, "bytes": size, "type": kind, "status": "changed" if rel in before else "created"}
        try:
            view = _view(path, kind, size)
        except Exception as exc:  # noqa: BLE001 - a file the agent left unreadable: said so, the answer goes on
            view = {"error": f"{type(exc).__name__}: {str(exc)[:300]}"}
        pictures, counted = view.pop("_pictures", None) or [], view.pop("_pictureCount", 0)
        text = str(view.pop("text", "") or "")
        keep = max(0, min(MAX_TEXT, text_left))
        if len(text) > keep:
            text, item["truncated"] = text[:keep], True
        text_left -= len(text)
        item["text"] = text
        item.update(view)
        if kind == "image":
            try:
                item["image"] = _picture(path.read_bytes(), budget) if size <= MAX_PARSE else {"error": f"larger than {MAX_PARSE // 1024 // 1024} MB: not read"}
            except OSError as exc:
                item["image"] = {"error": f"{type(exc).__name__}: {str(exc)[:200]}"}
        elif pictures:
            item["images"] = [{"name": name, **_picture(data, budget)} for name, data in pictures]
            if counted > len(pictures):
                item["imagesTotal"] = counted
        if size <= min(MAX_INLINE, inline_left):
            try:
                data = path.read_bytes()
                item.update(base64=base64.b64encode(data).decode("ascii"), sha256=hashlib.sha256(data).hexdigest())
                inline_left -= len(data)
            except OSError:
                pass
        out.append(item)
    return out, len(paths)


# -- the knowledge-base tool: an in-process MCP server the SDK bridges to Claude Code -------------------------------------
_KB_CLIENT = None
_KB_LOCK = threading.Lock()


def _kb_client():
    global _KB_CLIENT
    with _KB_LOCK:
        if _KB_CLIENT is None:
            _KB_CLIENT = boto3.client("bedrock-agent-runtime", region_name=REGION)
        return _KB_CLIENT


def _retrieve_text(question, kb_id):
    targets = [kb for kb in KNOWLEDGE_BASES if not kb_id or kb["id"] == kb_id]
    if not targets:
        raise ValueError(f"{kb_id} is not one of this agent's knowledge bases ({', '.join(kb['id'] for kb in KNOWLEDGE_BASES)})")
    parts = []
    for kb in targets:
        search = "vectorSearchConfiguration" if kb.get("type") == "VECTOR" else "managedSearchConfiguration"
        out = _kb_client().retrieve(knowledgeBaseId=kb["id"], retrievalQuery={"text": question[:1000]},
                                    retrievalConfiguration={search: {"numberOfResults": KB_TOP_K}})
        hits = []
        for index, found in enumerate(out.get("retrievalResults") or [], 1):
            text = ((found.get("content") or {}).get("text") or "").strip()
            where = found.get("location") or {}
            uri = ((where.get("s3Location") or {}).get("uri") or (where.get("webLocation") or {}).get("url")
                   or (where.get("customDocumentLocation") or {}).get("id") or "")
            score = found.get("score")
            head = f"[{index}] 出处 {uri.rsplit('/', 1)[-1] or '未知'}" + (f"，相关度 {score:.2f}" if isinstance(score, (int, float)) else "")
            hits.append(f"{head}\n{text}")
        parts.append(f"【{kb['name']}】\n" + ("\n\n".join(hits) if hits else "没有找到相关内容。"))
    return "\n\n".join(parts)


_KB_SCHEMA = {"type": "object", "properties": {"query": {"type": "string", "description": "要检索的问题或关键词（用用户的原话或其中的关键概念）"}},
              "required": ["query"]}
if len(KNOWLEDGE_BASES) > 1:
    _KB_SCHEMA["properties"]["kb_id"] = {"type": "string", "enum": [kb["id"] for kb in KNOWLEDGE_BASES], "description": "只检索这一个知识库；不填检索全部"}


@tool("retrieve", KB_DESCRIPTION, _KB_SCHEMA)
async def retrieve(args):
    question = str(args.get("query") or "").strip()
    if not question:
        return {"content": [{"type": "text", "text": "query 不能是空的"}], "is_error": True}
    try:
        text = await asyncio.to_thread(_retrieve_text, question, str(args.get("kb_id") or ""))
    except Exception as exc:  # noqa: BLE001 - an error result the model reads, the turn goes on
        logger.warning("retrieve failed: %s", exc)
        return {"content": [{"type": "text", "text": f"知识库检索失败：{type(exc).__name__}: {str(exc)[:300]}"}], "is_error": True}
    return {"content": [{"type": "text", "text": text}]}


# -- one turn ---------------------------------------------------------------------------------------------------------------
def _cli_line(line):
    line = str(line or "").rstrip()
    if line:
        logger.info("cli: %s", line[:500])


def _options(setup, *, resume=None, session_id=None, cwd=None, config=None, skills=None, sources=None, model=None):
    """The turn's options; a call with its own skills passes its directory, its config, its skill names and
    ``sources=["user"]``, the rest is the agent's own."""
    names = [s["name"] for s in SKILLS] if skills is None else list(skills)
    env = {"CLAUDE_CONFIG_DIR": str(config or setup["config"]), "CLAUDE_CODE_USE_BEDROCK": "1", "AWS_REGION": REGION,
           "CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC": "1", "DISABLE_AUTOUPDATER": "1", "CLAUDE_CODE_DISABLE_BUNDLED_SKILLS": "1",
           "CLAUDE_AGENT_SDK_DISABLE_BUILTIN_AGENTS": "1", "CLAUDE_CODE_DISABLE_AUTO_MEMORY": "1", "CLAUDE_AGENT_SDK_CLIENT_APP": "adlc-console-claude-agent/1"}
    if FAST_MODEL:
        env["ANTHROPIC_DEFAULT_HAIKU_MODEL"] = FAST_MODEL
    kwargs = {"system_prompt": SYSTEM_PROMPT, "model": model or MODEL, "max_turns": MAX_TURNS, "tools": list(BUILTIN_TOOLS) + (["Skill"] if names else []),
              "allowed_tools": list(BUILTIN_TOOLS) + ([KB_TOOL] if KNOWLEDGE_BASES else []), "skills": names,
              "mcp_servers": {KB_SERVER: create_sdk_mcp_server(name=KB_SERVER, tools=[retrieve])} if KNOWLEDGE_BASES else {},
              "strict_mcp_config": True, "permission_mode": "dontAsk", "cwd": str(cwd or setup["workspace"]), "env": env, "verbatim_prompts": True,
              "stderr": _cli_line}
    if sources is not None:
        kwargs["setting_sources"] = list(sources)
    if setup.get("cli"):
        kwargs["cli_path"] = setup["cli"]
    if EFFORT:
        kwargs["effort"] = EFFORT
    if MAX_BUDGET_USD:
        kwargs["max_budget_usd"] = MAX_BUDGET_USD
    if resume:
        kwargs["resume"] = resume
    elif session_id:
        kwargs["session_id"] = session_id
    return ClaudeAgentOptions(**kwargs)


def _brief(value, limit=200):
    text = json.dumps(value, ensure_ascii=False, default=str)
    return value if len(text) <= limit else text[:limit] + "…"


async def _run(prompt, options):
    calls, by_id, texts, result = [], {}, [], None
    async for message in query(prompt=prompt, options=options):
        if isinstance(message, AssistantMessage):
            for block in message.content:
                if isinstance(block, ToolUseBlock):
                    call = {"tool": block.name, "input": _brief(block.input), "ok": True}
                    if block.name == "Skill" and isinstance(block.input, dict):
                        call["skill"] = str(block.input.get("skill") or "")  # before _brief can cut a long input short
                    calls.append(call)
                    by_id[block.id] = call
                elif isinstance(block, TextBlock) and block.text.strip() and getattr(message, "parent_tool_use_id", None) is None:
                    texts.append(block.text.strip())
        elif isinstance(message, UserMessage) and isinstance(message.content, list):
            for block in message.content:
                if isinstance(block, ToolResultBlock) and block.is_error and block.tool_use_id in by_id:
                    by_id[block.tool_use_id]["ok"] = False
        elif isinstance(message, ResultMessage):
            result = message
    return calls, texts, result


async def _turn(runtime_session, prompt, plan=None, model=None):
    """One turn. A conversation's turn runs in the session's workspace and resumes its Claude Code session; a call
    of its own (``plan``: its skills, its input files) runs apart (its directory with its input files in it, its
    config, a session of its own, not resumed), the model told which files it was given."""
    setup = _setup()
    started = time.time()
    call = None
    if plan is not None:
        call = await asyncio.to_thread(_prepare_call, setup, plan)
        root, skip = call["workdir"], ()
        before = _manifest(root)  # the input files are in it: one the turn leaves alone is not the turn's
        claude, resumed = str(uuid.uuid4()), False
        asked = prompt + ("\n\n" + _inputs_note(call["inputs"]) if call["inputs"] else "")
        calls, texts, result = await _run(asked, _options(setup, session_id=claude, cwd=root, config=call["config"], skills=plan["names"],
                                                          sources=["user"], model=model))
    else:
        root, skip = setup["workspace"], ("calls",)
        before = _manifest(root, skip)
        claude = _conversations(setup).get(runtime_session) or str(uuid.uuid5(SESSION_NAMESPACE, runtime_session))
        resumed = _has_transcript(setup, claude)
        try:
            calls, texts, result = await _run(prompt, _options(setup, resume=claude, model=model) if resumed else
                                              _options(setup, session_id=claude, model=model))
        except Exception as exc:  # noqa: BLE001 - only a conversation that cannot go on is started again
            if claude not in str(exc):
                raise
            logger.warning("conversation %s cannot go on (%s): a new one starts", claude, str(exc)[:300])
            claude, resumed = str(uuid.uuid4()), False
            calls, texts, result = await _run(prompt, _options(setup, session_id=claude, model=model))
    if result is None:
        raise RuntimeError("Claude Code ended without a result")
    note = ""
    if result.is_error:
        if result.subtype != "error_max_turns" or not texts:
            raise RuntimeError(f"{result.subtype}: {result.result or '; '.join(result.errors or []) or 'no detail'}")
        note = f"\n\n〔用完了最多 {MAX_TURNS} 轮，回答可能不完整〕"
    claude = result.session_id or claude
    if call is None:
        _remember(setup, runtime_session, claude)
    artifacts, produced = await asyncio.to_thread(_artifacts, root, before, skip)
    answer = ((result.result if not result.is_error else "") or "\n\n".join(texts)).strip() + note
    out = {"result": answer, "tools": calls, "skills": [c["skill"] for c in calls if c.get("skill")],
           "skillsLoaded": list(dict.fromkeys(c["skill"] for c in calls if c.get("skill") and c.get("ok"))), "artifacts": artifacts,
           "seconds": round(time.time() - started, 2), "turns": result.num_turns, "costUsd": result.total_cost_usd, "usage": result.usage or {},
           "stopReason": result.stop_reason, "conversation": {"id": claude, "resumed": resumed}, "model": model or MODEL, "agent": AGENT.get("name"),
           "format": AGENT.get("format"), "workspace": {"path": str(setup["workspace"]), "sessionStorage": setup["persistent"]}}
    if produced > len(artifacts):
        out["artifactsTotal"] = produced
    if call is not None:
        out["call"] = {"id": call["id"], "workdir": call["workdir"].relative_to(setup["workspace"]).as_posix(), "skills": call["skills"],
                       "inputs": call["inputs"]}
    return out


def _label(call):
    if call["tool"] == KB_TOOL:
        return "retrieve"
    if call["tool"] == "Skill" and call.get("skill"):
        return f"Skill({call['skill']})"
    return call["tool"]


def _size(n):
    return f"{n} B" if n < 1024 else f"{n / 1024:.1f} KB" if n < 1024 * 1024 else f"{n / 1024 / 1024:.1f} MB"


def _footer(calls, artifacts=()):
    files = ("\n〔文件〕" + "，".join(f"{a['path']}（{_size(a['bytes'])}）" for a in artifacts)) if artifacts else ""
    if not calls:
        return "\n\n〔这一轮没有调用工具〕" + files
    counted = {}
    for call in calls:
        key = (_label(call), call.get("ok", True))
        counted[key] = counted.get(key, 0) + 1
    return "\n\n〔工具〕" + "，".join(f"{name}{'×' + str(n) if n > 1 else ''}{'' if ok else '（失败）'}" for (name, ok), n in counted.items()) + files


def _files(setup, limit=200):
    """The workspace's files (Claude Code's own directory left out), at most ``limit``."""
    root, out = setup["workspace"], []
    for folder, dirs, names in os.walk(root):
        dirs[:] = sorted(d for d in dirs if d != ".claude-agent")
        for name in sorted(names):
            path = Path(folder) / name
            try:
                stat_ = path.stat()
            except OSError:
                continue
            out.append({"path": str(path.relative_to(root)), "bytes": stat_.st_size, "modified": int(stat_.st_mtime)})
            if len(out) >= limit:
                return out
    return out


def _json(body, status=200):
    return Response(json.dumps(body, ensure_ascii=False, default=str), status_code=status, media_type="application/json")


@app.entrypoint
async def invoke(payload, context=None):
    """One turn: {"prompt": ...} → the reply as text/plain ({"format": "json"}: the turn's record, with the files it
    created or changed); with "skills" / "skillOverrides" / "files" → a call of its own, in its own directory with
    its input files; {"action": "workspace"} → the workspace's files and this session's conversation."""
    payload = payload if isinstance(payload, dict) else {"prompt": payload}
    runtime_session = str(getattr(context, "session_id", None) or payload.get("sessionId") or "local")
    if payload.get("action") == "workspace":
        setup = _setup()
        claude = _conversations(setup).get(runtime_session)
        return _json({"workspace": str(setup["workspace"]), "sessionStorage": setup["persistent"], "files": _files(setup),
                      "conversation": {"id": claude, "transcript": bool(claude) and _has_transcript(setup, claude)},
                      "skills": [s["name"] for s in SKILLS], "format": AGENT.get("format")})
    prompt = str(payload.get("prompt") or payload.get("message") or "").strip()
    if not prompt:
        return _json({"error": "the payload needs a non-empty prompt"}, 400)
    if len(prompt) > MAX_PROMPT:
        return _json({"error": f"the prompt is longer than {MAX_PROMPT} characters"}, 400)
    try:
        plan, model = _call_plan(payload), _call_model(payload)
    except PayloadError as exc:  # AgentCore drops a 4xx's body too: the log says why
        logger.warning("call refused %s", json.dumps({"session": runtime_session[:16], "status": 400, "error": str(exc)[:500]}, ensure_ascii=False))
        return _json({"error": str(exc)}, 400)
    lock = _LOCKS.setdefault(runtime_session, asyncio.Lock())
    async with lock:  # one turn at a time in a conversation
        try:
            out = await _turn(runtime_session, prompt, plan, model)
        except (PayloadError, SkillFetchError) as exc:
            status = 400 if isinstance(exc, PayloadError) else 503
            logger.warning("call refused %s", json.dumps({"session": runtime_session[:16], "status": status, "error": str(exc)[:500]}, ensure_ascii=False))
            return _json({"error": str(exc)}, status)
        except Exception as exc:  # noqa: BLE001 - one line to find in the log (AgentCore drops a 500's body), then the traceback
            failed = {"error": f"{type(exc).__name__}: {exc}"[:2000]}
            logger.error("turn failed %s", json.dumps({"session": runtime_session[:16], **failed}, ensure_ascii=False))
            logger.exception("traceback")
            return _json(failed, 500)
    logger.info("turn %s", json.dumps({"session": runtime_session[:16], "seconds": out["seconds"], "turns": out["turns"], "costUsd": out["costUsd"],
                                       "tools": [_label(c) for c in out["tools"]], "conversation": out["conversation"],
                                       "call": (out.get("call") or {}).get("id"), "skills": [s["name"] for s in (out.get("call") or {}).get("skills") or []],
                                       "inputs": [i["path"] for i in (out.get("call") or {}).get("inputs") or []],
                                       "files": [a["path"] for a in out["artifacts"]]}, ensure_ascii=False))
    if str(payload.get("format") or "").lower() == "json":
        return _json(out)
    trace = payload.get("trace")
    text = out["result"] + (_footer(out["tools"], out["artifacts"]) if trace is None or bool(trace) else "")
    return Response(text, media_type="text/plain; charset=utf-8")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "8080")))
'''


MARKER = re.compile(r"__(?:DOC|AGENT|MODEL|FAST_MODEL|MAX_TURNS|EFFORT|MAX_BUDGET_USD|BUILTIN_TOOLS|SKILLS|KNOWLEDGE_BASES|KB_TOP_K|KB_DESCRIPTION|"
                    r"SYSTEM_PROMPT|REGION)__")


def render_main(spec: Mapping[str, Any], kbs: Mapping[str, Mapping[str, Any]], skills: Mapping[str, Mapping[str, Any]], *, region: str,
                generated_at: str) -> str:
    doc = (f"{spec['name']} — a Claude Agent SDK agent generated by the ADLC console (Claude Agent SDK template).\n\n"
           f"Generated {generated_at} · model {spec['model']} · {len(spec['tools'])} built-in tool(s) · {len(spec['skills'])} skill(s) · "
           f"{len(spec['knowledgeBases'])} knowledge base(s)\n\n"
           "On AgentCore Runtime (a linux/arm64 container), POST /invocations with {\"prompt\": \"...\"} answers the reply as text/plain\n"
           "with a 〔工具〕 line (and a 〔文件〕 line for the files the turn made); {\"format\": \"json\"} answers the turn's record,\n"
           "its \"artifacts\" the files the turn created or changed (each with a text view); {\"action\": \"workspace\"} lists the\n"
           "workspace. A call may bring its own skills: \"skills\" (the names it may load, [] for none) and \"skillOverrides\"\n"
           "([{\"name\", \"uri\": \"s3://bucket/prefix/\"}], read from S3 for this call), and its input files: \"files\"\n"
           "([{\"path\", \"uri\": \"s3://bucket/key\", \"sha256\", \"bytes\"}], read into its directory before the prompt); such a\n"
           "call runs in its own directory calls/<id> of the workspace with its own Claude Code config, outside the session's\n"
           "conversation. An image the turn made, and the pictures of a document it made, come with a preview and a thumbnail.\n"
           "Claude Code runs on Bedrock with the runtime's execution role (CLAUDE_CODE_USE_BEDROCK=1). Locally:\n\n"
           "    docker build -t agent . && docker run -p 8080:8080 -e AWS_REGION -e AWS_ACCESS_KEY_ID -e AWS_SECRET_ACCESS_KEY -e AWS_SESSION_TOKEN agent\n"
           "    curl -s localhost:8080/invocations -H 'Content-Type: application/json' -d '{\"prompt\": \"你好\"}'\n\n"
           "Change the agent in the console rather than this file: every deployment generates it again.\n")
    agent = {"name": spec["name"], "description": spec["description"], "format": FORMAT, "generatedAt": generated_at, "sdk": SDK, "cli": CLAUDE_CODE}
    kb_literal = [{"id": k, "name": (kbs.get(k) or {}).get("name") or k, "type": (kbs.get(k) or {}).get("type") or "MANAGED"} for k in spec["knowledgeBases"]]
    skill_literal = [{"name": s["name"], "version": (skills.get(s["name"]) or {}).get("version")} for s in spec["skills"]]
    values = {
        "__DOC__": py_str(doc), "__AGENT__": py_value(agent), "__MODEL__": py_str(spec["model"]), "__FAST_MODEL__": py_str(spec["fastModel"] or ""),
        "__MAX_TURNS__": str(int(spec["maxTurns"])), "__EFFORT__": py_value(spec["effort"]), "__MAX_BUDGET_USD__": py_value(spec["maxBudgetUsd"]),
        "__BUILTIN_TOOLS__": py_value(list(spec["tools"])), "__SKILLS__": py_value(skill_literal), "__KNOWLEDGE_BASES__": py_value(kb_literal),
        "__KB_TOP_K__": str(int(spec["topK"])), "__KB_DESCRIPTION__": py_str(kb_tool_description(spec, kbs) if spec["knowledgeBases"] else "知识库检索"),
        "__SYSTEM_PROMPT__": py_str(system_prompt(spec, kbs)), "__REGION__": py_str(region),
    }
    # one pass over the template: a prompt or a name that happens to contain __MODEL__ is never replaced itself
    return MARKER.sub(lambda m: values[m.group(0)], MAIN_TEMPLATE)


def generate(spec: Mapping[str, Any], *, skills: Mapping[str, Mapping[str, Any]] | None = None, kbs: Mapping[str, Mapping[str, Any]] | None = None,
             region: str = "us-west-2", generated_at: str | None = None) -> dict[str, str]:
    """The build context of a checked spec: Dockerfile, main.py, requirements.txt, agent_spec.json and
    .claude/skills/<name>/SKILL.md for each skill (``skills``: name → {"text", "version"}; ``kbs``: id → {"name",
    "type", "description"} as AWS describes them)."""
    skills, kbs = skills or {}, kbs or {}
    missing = [s["name"] for s in spec["skills"] if s["name"] not in skills]
    if missing:
        raise ClaudeSdkError(f"技能 {'、'.join(missing)} 还没有读到它的 SKILL.md")
    when = generated_at or _now()
    main = render_main(spec, kbs, skills, region=region, generated_at=when)
    try:
        compile(main, "main.py", "exec")
    except SyntaxError as exc:  # pragma: no cover - a generator bug, never the user's
        raise ClaudeSdkError(f"生成的 main.py 编译不过（第 {exc.lineno} 行：{exc.msg}）：这是控制台的问题") from exc
    record = {"format": FORMAT, "generatedAt": when, "spec": dict(spec), "region": region,
              "skills": {n: {"version": s.get("version")} for n, s in skills.items() if n in {x["name"] for x in spec["skills"]}},
              "knowledgeBases": {k: v for k, v in kbs.items() if k in spec["knowledgeBases"]}, "versions": {"sdk": SDK, "cli": CLAUDE_CODE, "base": BASE_IMAGE}}
    files = {"Dockerfile": dockerfile(spec, when), "main.py": main, "requirements.txt": requirements(),
             "agent_spec.json": json.dumps(record, ensure_ascii=False, indent=1) + "\n"}
    for ref in spec["skills"]:
        files[f".claude/skills/{ref['name']}/SKILL.md"] = skills[ref["name"]]["text"]
    return files


def catalog() -> dict[str, Any]:
    """What the form offers: models, built-in tools, efforts, limits, the deploy options' defaults, the sample."""
    return {"models": [{"id": m, "label": label} for m, label in MODELS], "defaultModel": DEFAULT_MODEL, "defaultFastModel": DEFAULT_FAST_MODEL,
            "tools": [{"id": k, "label": v} for k, v in BUILTIN_TOOLS.items()], "efforts": list(EFFORTS),
            "limits": {"prompt": MAX_PROMPT, "skills": MAX_SKILLS, "knowledgeBases": MAX_KBS, "turns": [TURNS_MIN, TURNS_MAX], "topK": [TOP_K_MIN, TOP_K_MAX]},
            "defaults": {"maxTurns": DEFAULT_TURNS, "topK": DEFAULT_TOP_K, "systemPrompt": DEFAULT_PROMPT, "sample": DEFAULT_SAMPLE,
                         "mountPath": deploy.DEFAULT_MOUNT, "scanBlock": list(deploy.DEFAULT_SCAN_BLOCK),
                         "lifecycle": {"idleRuntimeSessionTimeout": 900, "maxLifetime": 28800}, "lifecycleRange": [deploy.LIFECYCLE_MIN, deploy.LIFECYCLE_MAX]},
            "severities": list(deploy.SEVERITIES), "sample": SAMPLE_SPEC, "kbTool": KB_TOOL,
            "versions": {"sdk": SDK, "cli": CLAUDE_CODE, "agentcore": AGENTCORE, "base": BASE_IMAGE, "node": NODE_MAJOR, "template": TEMPLATE}}


# -- the console's side ---------------------------------------------------------------------------------------------------

def build(console: Any, workspace: str, body: Mapping[str, Any]) -> dict[str, Any]:
    """A spec checked and generated against the workspace (its skills read from the library, its KBs described)."""
    spec = check_spec(body)
    ws, session = deploy._context(console, workspace)
    skills = load_skills(console, workspace, spec)
    kbs = resolve_kbs(session, ws["region"], spec["knowledgeBases"])
    files = generate(spec, skills=skills, kbs=kbs, region=ws["region"])
    return {"spec": spec, "files": files, "skills": skills, "kbs": kbs, "region": ws["region"]}


def preview(console: Any, workspace: str, body: Mapping[str, Any]) -> dict[str, Any]:
    built = build(console, workspace, body)
    data = bundle(built["files"])
    return {"spec": built["spec"], "files": built["files"], "size": len(data), "sha256": hashlib.sha256(data).hexdigest(),
            "skills": {n: s["version"] for n, s in built["skills"].items()}, "knowledgeBases": built["kbs"],
            "names": _names_view(built["spec"]["name"])}


def _names_view(name: str) -> dict[str, str]:
    names = names_for(name)
    return {"role": names.runtime_role(name), "repository": names.repository(name), "buildProject": names.build_project}


def download(console: Any, workspace: str, body: Mapping[str, Any]) -> dict[str, Any]:
    built = build(console, workspace, body)
    data = bundle(built["files"])
    return {"filename": f"{built['spec']['name']}-claude-agent.zip", "archive": base64.b64encode(data).decode("ascii"), "size": len(data),
            "files": sorted(built["files"])}


#: The deploy pipeline's options a template deployment passes through as they are.
DEPLOY_OPTIONS = ("endpoint", "smoke", "sessionStorage", "lifecycle", "network", "scanBlock")


def _key(workspace: str, name: str) -> str:
    return f"{workspace}:{name}"


def deploy_agent(console: Any, workspace: str, body: Mapping[str, Any], caller: Mapping[str, Any]) -> dict[str, Any]:
    """Generate the agent and start a deploy job (``deploy.start_deploy``, source dockerfile): a new runtime, or —
    ``runtimeId`` — a new version of one this console deployed. ``sessionStorage`` also tells the agent where its
    workspace is (``AGENT_WORKSPACE``)."""
    runtime_id = str(body.get("runtimeId") or "").strip() or None
    spec_body = dict(body)
    if runtime_id:
        record = deploy._records(console.store, workspace).get(runtime_id) or {}
        if not record.get("name"):
            raise ClaudeSdkError(f"{runtime_id} 不是这个控制台部署的 Runtime")
        spec_body["name"] = record["name"]
    built = build(console, workspace, spec_body)
    spec = built["spec"]
    ws, session = deploy._context(console, workspace)
    mount = None
    storage = deploy.check_session_storage(body.get("sessionStorage"))
    if storage:
        mount = storage[0]["sessionStorage"]["mountPath"]
    elif storage is None and runtime_id:  # a new version keeps the runtime's session storage: the agent must know where it is
        current = client(session, "bedrock-agentcore-control", ws["region"]).get_agent_runtime(agentRuntimeId=runtime_id)
        mount = next((f["sessionStorage"].get("mountPath") for f in current.get("filesystemConfigurations") or [] if "sessionStorage" in f), None)
    request: dict[str, Any] = {"source": "dockerfile", "archive": base64.b64encode(bundle(built["files"])).decode("ascii"),
                               "filename": f"{spec['name']}-claude-agent.zip", "protocol": "HTTP",
                               "environment": {"AGENT_WORKSPACE": mount.rstrip("/")} if mount else {},
                               "smokePrompt": spec["sample"] or DEFAULT_SAMPLE, "knowledgeBases": list(spec["knowledgeBases"])}
    for key in DEPLOY_OPTIONS:
        if body.get(key) is not None:
            request[key] = body[key]
    if runtime_id:
        request.pop("protocol")
    else:
        request["name"] = spec["name"]
    job = deploy.start_deploy(console, workspace, request, runtime_id=runtime_id, names=names_for(spec["name"]))
    entry = {"jobId": job["id"], "mode": "update" if runtime_id else "create", "runtimeId": runtime_id, "at": _now(), "by": caller.get("username"),
             "specHash": spec_hash(spec)[:16], "skills": {n: s["version"] for n, s in built["skills"].items()}, "sessionStorage": mount,
             "template": TEMPLATE, "model": spec["model"]}

    def change(all_: dict[str, Any]) -> dict[str, Any]:
        key = _key(workspace, spec["name"])
        rec = dict(all_.get(key) or {"name": spec["name"], "workspace": workspace, "createdAt": _now(), "createdBy": caller.get("username"),
                                     "deployments": []})
        rec.update(spec=spec, updatedAt=_now())
        rec["deployments"] = (list(rec.get("deployments") or []) + [entry])[-50:]
        return {**all_, key: rec}

    console.store.update(COLLECTION, {}, change)
    return {**job, "claude": entry}


def _view(console: Any, rec: Mapping[str, Any]) -> dict[str, Any]:
    deployments = []
    for d in reversed(rec.get("deployments") or []):
        job = studio._job(console, d.get("jobId")) or {}
        result, progress = job.get("result") or {}, job.get("progress") or {}
        deployments.append({**d, "runtimeId": d.get("runtimeId") or result.get("runtimeId") or progress.get("runtimeId"),
                            "version": result.get("version") or progress.get("version"), "jobStatus": job.get("status"), "jobError": job.get("error"),
                            "stage": progress.get("stage"), "scan": progress.get("scan"), "smoke": result.get("smoke")})
    runtime = next((d["runtimeId"] for d in deployments if d.get("runtimeId")), None)
    return {"name": rec.get("name"), "spec": rec.get("spec"), "createdAt": rec.get("createdAt"), "updatedAt": rec.get("updatedAt"),
            "createdBy": rec.get("createdBy"), "runtimeId": runtime, "deployments": deployments}


def list_agents(console: Any, workspace: str) -> list[dict[str, Any]]:
    mine = [r for r in (console.store.read(COLLECTION, {}) or {}).values() if r.get("workspace") == workspace]
    return [_view(console, r) for r in sorted(mine, key=lambda r: str(r.get("updatedAt") or ""), reverse=True)]


# -- routes ---------------------------------------------------------------------------------------------------------------

def _safe(fn: Any) -> Any:
    """Refusals answer 400, a missing runtime 404."""
    return safe(fn, (ClaudeSdkError, deploy.DeployError))


def register(router: Any) -> None:
    """``/workspaces/{wid}/claude-sdk/...``: the catalog; preview (the generated files) and bundle (the zip); deploy
    (admin: a new runtime, or a new version with ``runtimeId``); the agents made from the template."""
    add = router.add
    base = "/workspaces/{wid}/claude-sdk"
    add("GET", f"{base}/catalog", _safe(lambda r: (r.workspace(), (200, catalog()))[1]))
    add("POST", f"{base}/preview", _safe(lambda r: (200, preview(r.console, r.workspace(), r.body))))
    add("POST", f"{base}/bundle", _safe(lambda r: (200, download(r.console, r.workspace(), r.body))))
    add("POST", f"{base}/deploy", _safe(lambda r: (202, deploy_agent(r.console, r.workspace(), r.body, r.caller))), admin=True)
    add("GET", f"{base}/agents", _safe(lambda r: (200, {"agents": list_agents(r.console, r.workspace())})))
